"""PPT 原语：往母版新建自定义版式 + 按名定位版式。

PPT primitives: create a custom slide layout in the master, locate a layout by name.

提炼自参考脚本 ``gen_ppt.py``。PPT 的「带样式页面」靠版式(slideLayout)。python-pptx 只能
**用**已有版式、不能**新建**版式——新建必须做 OOXML 手术（把形状提升为 slideLayout 部件 +
双向关系 + 母版登记），是 PPT 模板最复杂的一环。:func:`create_custom_layout` 把这套「借壳换芯」
封装成一个带 ``decorate`` 回调的稳定原语：调用方只管在回调里画装饰，落户交给本原语。
"""

from __future__ import annotations

from collections.abc import Callable
from copy import deepcopy
from typing import Any

from lxml import etree
from pptx.opc.constants import CONTENT_TYPE, RELATIONSHIP_TYPE
from pptx.oxml.ns import qn
from pptx.parts.slide import SlideLayoutPart
from pptx.presentation import Presentation
from pptx.slide import SlideLayout

from ._ooxml import serialize_xml
from .errors import AnchorNotFoundError, ExternalMediaError

#: spTree 里必须保留、不属于「形状」的组属性元素（搬运形状时跳过它们）。
_GROUP_TAGS = (qn("p:nvGrpSpPr"), qn("p:grpSpPr"))
#: sldLayoutId 的 id 属性起始值（OOXML 约定 >= 0x80000000 区间）。
_SLD_LAYOUT_ID_BASE = 2147483649
#: OOXML 关系命名空间——形状里带此命名空间属性(r:embed/r:id/r:link)即引用了包关系(外部媒体)。
_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _check_base_index(layouts: Any, base_layout_index: int) -> None:
    """校验借壳版式索引在母版版式范围内，越界抛 :class:`ValueError`。"""
    if not 0 <= base_layout_index < len(layouts):
        raise ValueError(
            f"base_layout_index={base_layout_index} 越界；母版仅 {len(layouts)} 个版式 | out of range",
        )


def _remove_slide(prs: Presentation, index: int) -> None:
    """删除第 ``index`` 张幻灯片：移除其 ``sldId`` 登记并 drop 掉演示文稿→slide 的关系。"""
    slide_id_lst: Any = prs.slides._sldIdLst
    entry = list(slide_id_lst)[index]
    rel_id = entry.get(qn("r:id"))
    slide_id_lst.remove(entry)
    prs.part.drop_rel(rel_id)


def _first_external_rel(body: list[Any]) -> str | None:
    """扫描形状 XML 里对包关系的引用(``r:embed`` / ``r:id`` / ``r:link`` 等)；命中返回首个 rId，无则 None。

    这类引用只存在于**源部件**(slide)的 rels 里，深拷贝形状 XML 不会把 rels 一并搬来，会在新版式部件里悬空。
    """
    for shape in body:
        for el in shape.iter():
            for key, val in el.attrib.items():
                if key.startswith(f"{{{_REL_NS}}}"):
                    return str(val)
    return None


def _install_layout(prs: Presentation, name: str, body: list[Any], base_layout_index: int) -> SlideLayout:
    """把一组形状 ``body`` 落户成母版里名为 ``name`` 的新 slideLayout，返回该版式。

    「借壳换芯 + 落户」的共享核心：:func:`create_custom_layout`（形状来自 decorate 回调）与
    :func:`promote_slide_to_layout`（形状来自既有幻灯片）都复用它。全程走 lxml/opc——高层库无对应 API。

    ``body`` 含图片/超链接等包关系引用时抛 :class:`ExternalMediaError`（快速失败，避免静默产出损坏文件）。
    """
    dangling = _first_external_rel(body)
    if dangling is not None:
        raise ExternalMediaError(
            f"样板页/装饰形状引用了包关系(r:id={dangling!r}，如图片/超链接)，其 rels 不随形状搬运、会在新版式里悬空；"
            f"请改用不含外部媒体的样板页，或先移除这些媒体 | shape references a package relationship not carried over",
        )
    layouts = prs.slide_layouts
    _check_base_index(layouts, base_layout_index)  # 二次防线：promote 路径未在调用前早校验 base_index
    base_element: Any = layouts[base_layout_index].element

    # (c) 借 base 版式当合法外壳 -> 清空其形状、换上 body（借壳换芯）
    shell: Any = deepcopy(base_element)
    csld = shell.find(qn("p:cSld"))
    csld.set("name", name)
    shell.set("type", "blank")
    shell_sptree = csld.find(qn("p:spTree"))
    for child in list(shell_sptree):
        if child.tag not in _GROUP_TAGS:
            shell_sptree.remove(child)
    for child in body:
        shell_sptree.append(child)

    # (d) 落户：注册为 slideLayout 部件 + 双向关系 + 母版登记（高层库无 API，全靠 lxml/opc）
    master = prs.slide_masters[0]
    master_part: Any = master.part
    package: Any = prs.part.package
    partname = package.next_partname("/ppt/slideLayouts/slideLayout%d.xml")
    layout_part = SlideLayoutPart.load(partname, CONTENT_TYPE.PML_SLIDE_LAYOUT, package, serialize_xml(shell))
    layout_part.relate_to(master_part, RELATIONSHIP_TYPE.SLIDE_MASTER)  # layout -> master（规范必需）
    rel_id = master_part.relate_to(layout_part, RELATIONSHIP_TYPE.SLIDE_LAYOUT)  # master -> layout，拿 rId
    master_element: Any = master.element
    id_lst = master_element.find(qn("p:sldLayoutIdLst"))
    existing_ids = [int(e.get("id")) for e in id_lst.findall(qn("p:sldLayoutId"))]
    entry = etree.SubElement(id_lst, qn("p:sldLayoutId"))
    entry.set("id", str(max(existing_ids) + 1 if existing_ids else _SLD_LAYOUT_ID_BASE))
    entry.set(qn("r:id"), rel_id)

    return locate_by_master(prs, name)


def create_custom_layout(
    prs: Presentation,
    name: str,
    decorate: Callable[[Any], None],
    *,
    base_layout_index: int = 6,
) -> SlideLayout:
    """在母版里新建一个名为 ``name`` 的自定义版式，装饰由 ``decorate`` 回调绘制，返回新版式。

    Create a custom slide layout named ``name`` in the first master.

    - ``decorate(shapes)`` 收到一个**临时幻灯片**的 ``shapes``（python-pptx ``SlideShapes``），
      在其上用常规 python-pptx API 画装饰（矩形/文本框/竖条等）——因为 python-pptx 只能往 slide
      画、不能往 layout 画。本原语随后把这些形状「偷」进新版式，并删除临时页。
    - ``base_layout_index`` 指定借作合法外壳的既有版式（默认 6 = 空白版式）。**该 base 版式应不含
      外部关系**（背景图片 / 超链接等）：新版式部件只补建 layout→master 一条关系，深拷贝 base XML 里
      对媒体/hlink 的 ``r:id`` 引用不会随之搬运，会在新部件里悬空。默认空白版式无此问题；若必须基于
      含媒体的版式，请自行搬运其 rels。
    - 落户全过程走 lxml/opc：注册为 slideLayout 部件 + layout↔master 双向关系 + 母版 sldLayoutIdLst 登记。
    """
    layouts = prs.slide_layouts
    _check_base_index(layouts, base_layout_index)  # 早校验：坏索引不应先污染 deck（加临时页）
    base = layouts[base_layout_index]

    # (a) 借临时 slide 画版式装饰（python-pptx 只能往 slide 画）
    tmp = prs.slides.add_slide(base)
    tmp_sptree: Any = tmp.shapes._spTree
    for child in list(tmp_sptree):
        if child.tag not in _GROUP_TAGS:
            tmp_sptree.remove(child)
    decorate(tmp.shapes)

    # (b) 偷形状 + 删临时 slide（否则成品会多一页）
    body = [deepcopy(child) for child in tmp_sptree if child.tag not in _GROUP_TAGS]
    _remove_slide(prs, len(prs.slides) - 1)

    return _install_layout(prs, name, body, base_layout_index)


def promote_slide_to_layout(
    prs: Presentation,
    slide_index: int,
    name: str,
    *,
    base_layout_index: int = 6,
    remove_source: bool = False,
) -> SlideLayout:
    """把第 ``slide_index`` 张**既有幻灯片**升级为母版里名为 ``name`` 的可复用版式，返回该版式。

    Promote an existing slide into a reusable named slide layout in the master.

    这是 W3 抽模板的 PPT 方向：参考 deck 里有一张精心设计的样板页，把它**脱胎**为母版版式后，
    后续 :func:`locate_by_master` + ``add_slide`` 即可按该版式批量出图、保持品牌一致——与
    :func:`create_custom_layout`（从回调新建）互为「既有页 → 版式」与「无中生有 → 版式」两条入口，
    共用同一落户核心。

    - ``slide_index`` 为源幻灯片在 ``prs.slides`` 中的序号（越界抛 :class:`ValueError`）。
    - 源页形状被**深拷贝**进新版式，源页默认保留（``remove_source=True`` 则一并删除源页，产出更干净的模板）。
    - ``base_layout_index`` 同 :func:`create_custom_layout`：借作合法外壳的既有版式，应不含外部关系。
    - **不搬运源页的外部关系**：图片/超链接等媒体的 ``r:id`` 只存在于源 slide 部件，深拷贝形状 XML 后会在
      新版式部件里悬空。故源页含此类媒体时**快速失败**抛 :class:`ExternalMediaError`（避免静默产出损坏文件）——
      请改用不含外部媒体的样板页，或先移除这些媒体。
    """
    slides = prs.slides
    count = len(slides)
    if not 0 <= slide_index < count:
        raise ValueError(f"slide_index={slide_index} 越界；共 {count} 张幻灯片 | slide index out of range")
    slide = slides[slide_index]
    src_sptree: Any = slide.shapes._spTree
    body = [deepcopy(child) for child in src_sptree if child.tag not in _GROUP_TAGS]

    layout = _install_layout(prs, name, body, base_layout_index)
    if remove_source:
        _remove_slide(prs, slide_index)
    return layout


def locate_by_master(prs: Presentation, name: str) -> SlideLayout:
    """在所有母版中按名定位版式(slideLayout)，返回匹配的 :class:`SlideLayout`。

    Locate a slide layout by ``name`` across every master. 找不到抛 :class:`AnchorNotFoundError`。
    W3 抽模板 / W1 复用版式时据此按名取版式（版式名 = ``cSld/@name``）。
    """
    for master in prs.slide_masters:
        for layout in master.slide_layouts:
            if layout.name == name:
                found: SlideLayout = layout
                return found
    raise AnchorNotFoundError(f"版式未找到 | slide layout not found: {name!r}")
