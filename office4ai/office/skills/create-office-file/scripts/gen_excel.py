"""参考脚本：从零创建 Excel 工作簿 | Reference: create an Excel workbook from scratch.

演示：表头 + 数据行 + 公式单元格 + 柱状图 + 命名区域。产物存到 cwd。
注意 openpyxl **不计算公式值**（只写公式字符串），真正的值由 Office/LibreOffice 打开时重算。
"""

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName

wb = Workbook()
ws = wb.active
ws.title = "销售数据"

# 表头（加粗 + 底色）
header_fill = PatternFill("solid", fgColor="1F4E79")
for col, name in enumerate(["月份", "销售额", "成本"], start=1):
    cell = ws.cell(1, col, name)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = header_fill

# 数据行
data = [("1月", 120, 80), ("2月", 150, 95), ("3月", 135, 88), ("4月", 180, 110)]
for r, (month, sales, cost) in enumerate(data, start=2):
    ws.cell(r, 1, month)
    ws.cell(r, 2, sales)
    ws.cell(r, 3, cost)

# 公式：毛利合计（openpyxl 只写公式，值待重算）
ws["E1"] = "毛利合计"
ws["E2"] = "=SUM(B2:B5)-SUM(C2:C5)"

# 柱状图（销售额 vs 月份）
chart = BarChart()
chart.title = "月度销售额"
chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=5), titles_from_data=True)
chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=5))
ws.add_chart(chart, "G1")

# 命名区域（供下游/模板引用）
wb.defined_names["销售额区"] = DefinedName("销售额区", attr_text="销售数据!$B$2:$B$5")

wb.save("sales_book.xlsx")
print("created sales_book.xlsx")
