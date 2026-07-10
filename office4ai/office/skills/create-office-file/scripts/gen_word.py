"""参考脚本：从零创建「红头文件」Word 文档 | Reference: create a red-header Word doc from scratch.

读之学模式后，写你自己的几行 glue 提交给 ``office_run_script``。脚本在沙箱里运行，产物存到
当前工作目录（cwd == office_run_script 的 work_dir）。**你决定文件名与所有动作。**

红头文件的招牌 = 红色居中加粗的发文机关名 + 一条红线 + 文号 + 标题 + 正文 + 落款。
python-docx 对「段落下红线」这类边框没有直接 API，需下探到 OOXML（本脚本演示这一常见坑）。
"""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

RED = RGBColor(0xC0, 0x00, 0x00)


def _add_red_bottom_border(paragraph) -> None:
    """给段落加一条红色下边框（红头文件的「红线」）——python-docx 无直接 API，下探 OOXML。"""
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "18")  # 边框粗细（1/8 pt）
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "C00000")
    borders.append(bottom)
    p_pr.append(borders)


doc = Document()

# 发文机关（红色、居中、加粗、大字号）
org = doc.add_paragraph()
org.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = org.add_run("XX 市人民政府办公室")
run.font.color.rgb = RED
run.font.bold = True
run.font.size = Pt(30)
_add_red_bottom_border(org)

# 文号（居中，红线之下）
wenhao = doc.add_paragraph()
wenhao.alignment = WD_ALIGN_PARAGRAPH.CENTER
wenhao.add_run("XX 府办发〔2026〕1 号").font.size = Pt(14)

# 标题
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
t_run = title.add_run("关于开展 2026 年度工作的通知")
t_run.font.bold = True
t_run.font.size = Pt(20)

# 正文
doc.add_paragraph("各有关单位：").runs[0].font.size = Pt(14)
body = doc.add_paragraph(
    "为做好 2026 年度各项工作，经研究，现就有关事项通知如下。请遵照执行。",
)
body.paragraph_format.first_line_indent = Pt(28)
body.runs[0].font.size = Pt(14)

# 落款（右对齐）
signoff = doc.add_paragraph()
signoff.alignment = WD_ALIGN_PARAGRAPH.RIGHT
signoff.add_run("XX 市人民政府办公室\n2026 年 7 月 10 日").font.size = Pt(14)

doc.save("red_header.docx")
print("created red_header.docx")
