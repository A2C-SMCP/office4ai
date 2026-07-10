"""参考脚本：复用既有 .pptx 的母版版式加页 | Reference: reuse a slide master layout.

母版样式复用：从既有品牌 deck 里按名取母版版式，用它 ``add_slide`` 加新页，保持全 deck 品牌一致，
而不是新建一张裸页再手动仿样式。``locate_by_master(prs, 版式名)`` 跨所有母版按名（``cSld/@name``）找版式。

⚠️ 填路径示意：``TARGET`` 换成你**本地**的真实 deck 路径（禁网）；``LAYOUT_NAME`` 换成该 deck 里
   真实存在的版式名（找不到抛 AnchorNotFoundError）。
"""

from pptx import Presentation

from office4ai.office.authoring.helpers import locate_by_master

TARGET = "/abs/local/path/to/branded_deck.pptx"
LAYOUT_NAME = "标题和内容"

prs = Presentation(TARGET)  # 读既有 deck（禁网：本地路径）
layout = locate_by_master(prs, LAYOUT_NAME)  # 按名取母版版式
slide = prs.slides.add_slide(layout)  # 复用该版式加新页
if slide.shapes.title is not None:
    slide.shapes.title.text = "新增章节"
prs.save("deck_extended.pptx")
print(f"reused master layout {LAYOUT_NAME!r} -> deck_extended.pptx")
