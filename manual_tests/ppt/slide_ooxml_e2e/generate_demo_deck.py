"""High-fidelity slide generator for the whole-slide OOXML path (office4ai #40).

Demonstrates the styling headroom that OOXML unlocks versus the fine-grained
Office.js tools: rounded cards with fills, an accent header bar, a styled table,
and a NATIVE chart — all baked into a .pptx package the Server can hand to
``ppt_insert_slides_ooxml`` via ``source_path``.

Run::

    uv run python manual_tests/ppt/slide_ooxml_e2e/generate_demo_deck.py [out.pptx]

The generated file is then passed as ``source_path`` to ppt_insert_slides_ooxml
(live PowerPoint + Add-In required for the actual insert step).
"""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Pt

# 16:9 canvas
SLIDE_W = Emu(12192000)
SLIDE_H = Emu(6858000)

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
BLUE = RGBColor(0x2E, 0x6F, 0xB5)
LIGHT = RGBColor(0xF2, 0xF4, 0xF7)
GREY = RGBColor(0x6B, 0x72, 0x80)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def _card(slide, left, top, width, height, title, value, caption):
    """A rounded KPI card: rounded rectangle fill + title/value/caption text."""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = LIGHT
    shape.line.fill.background()
    shape.shadow.inherit = False

    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(12)
    tf.margin_top = Pt(8)
    p0 = tf.paragraphs[0]
    p0.text = title
    p0.font.size = Pt(11)
    p0.font.color.rgb = GREY

    p1 = tf.add_paragraph()
    p1.text = value
    p1.font.size = Pt(24)
    p1.font.bold = True
    p1.font.color.rgb = NAVY

    p2 = tf.add_paragraph()
    p2.text = caption
    p2.font.size = Pt(9)
    p2.font.color.rgb = GREY


def build_deck(out_path: Path) -> Path:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

    # ── accent header bar ──
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Pt(54))
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    bar.shadow.inherit = False
    tb = bar.text_frame
    tb.margin_left = Pt(28)
    tp = tb.paragraphs[0]
    tp.text = "2024 年度人员成本分析报告"
    tp.font.size = Pt(20)
    tp.font.bold = True
    tp.font.color.rgb = WHITE

    # ── KPI cards ──
    card_w, card_h, gap, x0, y0 = Pt(150), Pt(70), Pt(14), Pt(28), Pt(72)
    data = [
        ("总成本", "¥1,280万", "含工资·社保·公积金"),
        ("人均成本", "¥18.6万", "在岗 69 人"),
        ("同比增长", "+8.5%", "较 2023 年度"),
        ("预算执行率", "92.3%", "预算 ¥1,387万"),
    ]
    for i, (t, v, c) in enumerate(data):
        _card(slide, Emu(int(x0) + i * (int(card_w) + int(gap))), y0, card_w, card_h, t, v, c)

    # ── styled table ──
    rows, cols = 4, 4
    tbl_shape = slide.shapes.add_table(rows, cols, x0, Pt(160), Pt(330), Pt(150))
    table = tbl_shape.table
    headers = ["部门", "人数", "总成本(万)", "占比"]
    body = [
        ["技术研发部", "22", "520", "40.6%"],
        ["市场销售部", "15", "285", "22.3%"],
        ["运营管理部", "12", "198", "15.5%"],
    ]
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = BLUE
        para = cell.text_frame.paragraphs[0]
        para.font.color.rgb = WHITE
        para.font.bold = True
        para.font.size = Pt(11)
    for r, record in enumerate(body, start=1):
        for j, val in enumerate(record):
            cell = table.cell(r, j)
            cell.text = val
            cell.text_frame.paragraphs[0].font.size = Pt(11)

    # ── native chart ──
    chart_data = CategoryChartData()
    chart_data.categories = ["技术研发", "市场销售", "运营管理", "财务行政"]
    chart_data.add_series("总成本(万元)", (520, 285, 198, 152))
    gframe = slide.shapes.add_chart(
        XL_CHART_TYPE.BAR_CLUSTERED, Pt(380), Pt(160), Pt(300), Pt(150), chart_data
    )
    chart = gframe.chart
    chart.has_legend = False
    plot = chart.plots[0]
    plot.has_data_labels = True
    plot.data_labels.font.size = Pt(9)

    # ── footer insight ──
    foot = slide.shapes.add_textbox(x0, Pt(320), Pt(660), Pt(40))
    fp = foot.text_frame.paragraphs[0]
    fp.text = "关键发现：技术研发部成本占比最高 (40.6%)，人均成本达 23.6 万，建议优化人员结构配置。"
    fp.font.size = Pt(10)
    fp.font.color.rgb = GREY
    fp.alignment = PP_ALIGN.LEFT

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return out_path


if __name__ == "__main__":
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("/tmp/office4ai_ooxml_demo.pptx")
    saved = build_deck(out)
    size = saved.stat().st_size
    print(f"✅ generated {saved} ({size} bytes, base64 ≈ {int(size * 4 / 3)} chars)")
