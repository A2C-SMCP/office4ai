"""range_e2e 夹具构建器。

生成单个多 sheet 夹具 ``grid.xlsx``（提交入库，供真机打开），三张表覆盖 #30 全部
range 用例：

- ``Data`` —— A1:C4 预填数值网格（Region/Q1/Q2 + 3 行），activeSheet。用于
  ``get:range`` 读已有数据、``copy:range`` 源数据、``set:formula`` 的数值参与计算。
- ``Blank`` —— 空表。用于 ``set:range`` 写入与 ``clear:range``（先填后清）不污染 Data。
- ``Shift`` —— 3×3 自描述标签网格（每格值=自身地址 "A1".."C3"）。用于
  ``insert:range`` / ``delete:range`` 验证单元格位移后内容落到预期位置。

重新生成：
    uv run python manual_tests/excel/range_e2e/_fixtures.py
"""

from __future__ import annotations

from pathlib import Path

FIXTURE_DIR = Path(__file__).parent.parent / "fixtures" / "range_e2e"

# Data 表：A1:C4 数值网格（Q1/Q2 为数字，供 set:formula 的 SUM/算术引用）。
_DATA_ROWS: list[list[object]] = [
    ["Region", "Q1", "Q2"],
    ["North", 100, 150],
    ["South", 200, 250],
    ["East", 300, 350],
]

# Shift 表：3×3 自描述网格，每格值等于自身 A1 地址，便于断言位移后「谁落到哪」。
_SHIFT_ROWS: list[list[object]] = [
    ["A1", "B1", "C1"],
    ["A2", "B2", "C2"],
    ["A3", "B3", "C3"],
]


def _fill(ws, rows: list[list[object]]) -> None:
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=value)


def build_grid(path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    data = wb.active
    data.title = "Data"
    _fill(data, _DATA_ROWS)
    wb.create_sheet("Blank")  # 空表，供 set/clear 写入不污染 Data
    shift = wb.create_sheet("Shift")
    _fill(shift, _SHIFT_ROWS)
    wb.active = 0  # activeSheet = Data
    wb.save(str(path))


def build_all() -> dict[str, Path]:
    FIXTURE_DIR.mkdir(parents=True, exist_ok=True)
    builders = {"grid.xlsx": build_grid}
    paths: dict[str, Path] = {}
    for name, builder in builders.items():
        target = FIXTURE_DIR / name
        builder(target)
        paths[name] = target
        print(f"📝 生成夹具: {target}")
    return paths


def ensure_fixtures() -> dict[str, Path]:
    """确保 grid.xlsx 夹具存在（缺失才生成）。"""
    FIXTURE_DIR.mkdir(parents=True, exist_ok=True)
    if (FIXTURE_DIR / "grid.xlsx").exists():
        return {"grid.xlsx": FIXTURE_DIR / "grid.xlsx"}
    return build_all()


if __name__ == "__main__":
    build_all()
