"""format_e2e 夹具构建器。

生成单个多 sheet 夹具 ``fmt.xlsx``（提交入库，供真机打开）：

- ``Data`` —— A1:C4 数值网格（Region/Q1/Q2 + 3 行），activeSheet。用于
  set/get:rangeFormat（字体/填充/对齐/边框/数字格式）与 add/clear:conditionalFormat
  （B 列数值供 cellValue 规则）。
- ``Merge`` —— A1:C1 = M1/M2/M3 标签行。用于 merge/unmerge:cells，验证合并保留左上值、
  其余清空，以及取消合并后 merged_ranges 清空。

重新生成：
    uv run python manual_tests/excel/format_e2e/_fixtures.py
"""

from __future__ import annotations

from pathlib import Path

FIXTURE_DIR = Path(__file__).parent.parent / "fixtures" / "format_e2e"

# Data 表：A1:C4 数值网格（Q1/Q2 为数字，供 numberFormat / 条件格式）。
_DATA_ROWS: list[list[object]] = [
    ["Region", "Q1", "Q2"],
    ["North", 100, 150],
    ["South", 200, 250],
    ["East", 300, 350],
]

# Merge 表：A1:C1 标签行，合并后只保留左上 M1。
_MERGE_ROWS: list[list[object]] = [
    ["M1", "M2", "M3"],
]


def _fill(ws, rows: list[list[object]]) -> None:
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=value)


def build_fmt(path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    data = wb.active
    data.title = "Data"
    _fill(data, _DATA_ROWS)
    merge = wb.create_sheet("Merge")
    _fill(merge, _MERGE_ROWS)
    wb.active = 0  # activeSheet = Data
    wb.save(str(path))


def build_all() -> dict[str, Path]:
    FIXTURE_DIR.mkdir(parents=True, exist_ok=True)
    builders = {"fmt.xlsx": build_fmt}
    paths: dict[str, Path] = {}
    for name, builder in builders.items():
        target = FIXTURE_DIR / name
        builder(target)
        paths[name] = target
        print(f"📝 生成夹具: {target}")
    return paths


def ensure_fixtures() -> dict[str, Path]:
    """确保 fmt.xlsx 夹具存在（缺失才生成）。"""
    FIXTURE_DIR.mkdir(parents=True, exist_ok=True)
    if (FIXTURE_DIR / "fmt.xlsx").exists():
        return {"fmt.xlsx": FIXTURE_DIR / "fmt.xlsx"}
    return build_all()


if __name__ == "__main__":
    build_all()
