"""read_state_e2e 夹具构建器。

生成三个 .xlsx 夹具（提交入库，供真机打开）：

- ``multi_sheet.xlsx`` —— 多 sheet（Sheet1 空 / Data 预填 / Report 空），用于
  workbookInfo 多表、activeSheet、fileName 与 worksheetInfo 用例。
- ``hidden_sheet.xlsx`` —— 含隐藏 sheet（Visible 可见 / Hidden 隐藏），用于
  workbookInfo 的 isHidden 用例。
- ``prefilled.xlsx`` —— 单 Data 表，预填混合类型（字符串/数字/布尔/0/False/空），
  用于 selectedRange 的 2D values 与混合类型用例（真机需先手动选区）。

重新生成：
    uv run python manual_tests/excel/read_state_e2e/_fixtures.py
"""

from __future__ import annotations

from pathlib import Path

FIXTURE_DIR = Path(__file__).parent.parent / "fixtures" / "read_state_e2e"

# Data 表预填内容：A1:C4，usedRange = 4 行 × 3 列
_DATA_ROWS: list[list[object]] = [
    ["Region", "Q1", "Q2"],
    ["North", 100, 150],
    ["South", 200, 250],
    ["East", 300, 350],
]

# prefilled 混合类型：A1:C2（字符串/数字/布尔 + falsy 0/False/空）
_MIXED_ROWS: list[list[object]] = [
    ["Hello", 42, True],
    [0, False, ""],
]


def _fill(ws, rows: list[list[object]]) -> None:
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=value)


def build_multi_sheet(path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    sheet1 = wb.active
    sheet1.title = "Sheet1"  # 空表，验证空 usedRange
    data = wb.create_sheet("Data")
    _fill(data, _DATA_ROWS)
    wb.create_sheet("Report")  # 空表
    # 激活表设为 Sheet1（index 0）
    wb.active = 0
    wb.save(str(path))


def build_hidden_sheet(path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    visible = wb.active
    visible.title = "Visible"
    _fill(visible, _DATA_ROWS)
    hidden = wb.create_sheet("Hidden")
    _fill(hidden, [["secret", 1]])
    hidden.sheet_state = "hidden"
    wb.active = 0
    wb.save(str(path))


def build_prefilled(path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    data = wb.active
    data.title = "Data"
    _fill(data, _MIXED_ROWS)
    wb.save(str(path))


def build_all() -> dict[str, Path]:
    FIXTURE_DIR.mkdir(parents=True, exist_ok=True)
    builders = {
        "multi_sheet.xlsx": build_multi_sheet,
        "hidden_sheet.xlsx": build_hidden_sheet,
        "prefilled.xlsx": build_prefilled,
    }
    paths: dict[str, Path] = {}
    for name, builder in builders.items():
        target = FIXTURE_DIR / name
        builder(target)
        paths[name] = target
        print(f"📝 生成夹具: {target}")
    return paths


def ensure_fixtures() -> dict[str, Path]:
    """确保三个夹具存在（缺失才生成），返回名→路径。"""
    FIXTURE_DIR.mkdir(parents=True, exist_ok=True)
    if all((FIXTURE_DIR / n).exists() for n in ("multi_sheet.xlsx", "hidden_sheet.xlsx", "prefilled.xlsx")):
        return {n: FIXTURE_DIR / n for n in ("multi_sheet.xlsx", "hidden_sheet.xlsx", "prefilled.xlsx")}
    return build_all()


if __name__ == "__main__":
    build_all()
