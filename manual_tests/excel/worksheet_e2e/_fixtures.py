"""worksheet_e2e 夹具构建器。

生成单个多 sheet 夹具 ``book.xlsx``（提交入库，供真机打开），覆盖 #32 全部 worksheet
管理用例：增（add）/ 删（delete）/ 改名（rename）/ 激活（activate）/ 列举（get:worksheets）。

四张自描述表（每格 A1 值即「<表名>-A1」，便于断言删除/改名后哪张表还在）：

- ``Alpha`` —— **activeSheet**，A1:B2 小网格。get:worksheets 断言它 isActive。
- ``Beta``  —— 改名 / 删除目标。
- ``Gamma`` —— 激活目标。
- ``Delta`` —— 隔离表：删除/改名其它表时它须原样保留。

每个 case 都会从夹具复制一份全新工作副本再打开，故各 case 间的增删改互不污染。

重新生成：
    uv run python manual_tests/excel/worksheet_e2e/_fixtures.py
"""

from __future__ import annotations

from pathlib import Path

FIXTURE_DIR = Path(__file__).parent.parent / "fixtures" / "worksheet_e2e"

# 表名 → 该表 A1:B2 内容（自描述，A1 = "<表名>-A1"，便于读盘核对身份）。
_SHEETS: dict[str, list[list[object]]] = {
    "Alpha": [["Alpha-A1", "active"], ["row2", "data"]],
    "Beta": [["Beta-A1", "rename/delete target"]],
    "Gamma": [["Gamma-A1", "activate target"]],
    "Delta": [["Delta-A1", "isolation guard"]],
}


def _fill(ws, rows: list[list[object]]) -> None:
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=value)


def build_book(path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    first = True
    for name, rows in _SHEETS.items():
        if first:
            ws = wb.active
            ws.title = name
            first = False
        else:
            ws = wb.create_sheet(name)
        _fill(ws, rows)
    wb.active = 0  # activeSheet = Alpha
    wb.save(str(path))


def build_all() -> dict[str, Path]:
    FIXTURE_DIR.mkdir(parents=True, exist_ok=True)
    paths: dict[str, Path] = {}
    for name, builder in {"book.xlsx": build_book}.items():
        target = FIXTURE_DIR / name
        builder(target)
        paths[name] = target
        print(f"📝 生成夹具: {target}")
    return paths


def ensure_fixtures() -> dict[str, Path]:
    """确保 book.xlsx 夹具存在（缺失才生成）。"""
    FIXTURE_DIR.mkdir(parents=True, exist_ok=True)
    if (FIXTURE_DIR / "book.xlsx").exists():
        return {"book.xlsx": FIXTURE_DIR / "book.xlsx"}
    return build_all()


if __name__ == "__main__":
    build_all()
