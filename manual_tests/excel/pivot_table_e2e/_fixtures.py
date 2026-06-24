"""pivot_table_e2e 夹具构建器。

生成单个夹具 ``pivot.xlsx``（提交入库，供真机打开），覆盖 #35 全部 pivotTable 用例：
建透视表（insert）/ 列举（get:pivotTables）/ 删除（delete）。

两张表：

- ``Data`` —— **activeSheet**，A1:C5 数据源（Region/Product/Amount + 4 行）。源区 A1:C5，
  目标落点用同表空白区 E1 / E20（source 与 target 须同一 worksheet——AddIn 两个 getRange
  都基于同一 worksheet）。
- ``Blank`` —— 空表。供 ``get:pivotTables`` 验证「无透视表 → 空列表」。

透视表为视觉对象，openpyxl 对 Excel 透视表回读支持有限，故本套件以**协议层
get:pivotTables 回读**为唯一验证依据（不构造 rows/columns/values/filters，仅
sourceAddress + targetAddress 创建空透视表骨架）。

重新生成：
    uv run python manual_tests/excel/pivot_table_e2e/_fixtures.py
"""

from __future__ import annotations

from pathlib import Path

FIXTURE_DIR = Path(__file__).parent.parent / "fixtures" / "pivot_table_e2e"

# Data：透视表数据源（含表头行 + 可聚合的数值列 Amount）。
_DATA_ROWS: list[list[object]] = [
    ["Region", "Product", "Amount"],
    ["East", "A", 100],
    ["East", "B", 150],
    ["West", "A", 200],
    ["West", "B", 250],
]


def _fill(ws, rows: list[list[object]]) -> None:
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=value)


def build_pivot(path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    data = wb.active
    data.title = "Data"
    _fill(data, _DATA_ROWS)
    wb.create_sheet("Blank")  # 空表，供 get:pivotTables 验证空列表
    wb.active = 0  # activeSheet = Data
    wb.save(str(path))


def build_all() -> dict[str, Path]:
    FIXTURE_DIR.mkdir(parents=True, exist_ok=True)
    paths: dict[str, Path] = {}
    for name, builder in {"pivot.xlsx": build_pivot}.items():
        target = FIXTURE_DIR / name
        builder(target)
        paths[name] = target
        print(f"📝 生成夹具: {target}")
    return paths


def ensure_fixtures() -> dict[str, Path]:
    """确保 pivot.xlsx 夹具存在（缺失才生成）。"""
    FIXTURE_DIR.mkdir(parents=True, exist_ok=True)
    if (FIXTURE_DIR / "pivot.xlsx").exists():
        return {"pivot.xlsx": FIXTURE_DIR / "pivot.xlsx"}
    return build_all()


if __name__ == "__main__":
    build_all()
