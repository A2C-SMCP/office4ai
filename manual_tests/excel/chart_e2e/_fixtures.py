"""chart_e2e 夹具构建器。

生成单个夹具 ``chart.xlsx``（提交入库，供真机打开），覆盖 #34 全部 chart 用例：
建图（insert）/ 列举（get:charts）/ 更新（update）/ 删除（delete）。

两张表：

- ``Data`` —— **activeSheet**，A1:C4 数值网格（类别 + Sales/Cost 两列数值）。作所有图表的
  数据源（``sourceAddress`` 多用 ``A1:C4`` 或 ``A1:B4`` 单列）。
- ``Blank`` —— 空表。供 ``get:charts`` 验证「无图表 → 空列表」。

图表为视觉对象，openpyxl 对 Excel 原生图表的回读支持有限（``chart_count`` 仅 best-effort），
故本套件以**协议层 get:charts 回读**为主验证，openpyxl chart_count 仅作旁证。

重新生成：
    uv run python manual_tests/excel/chart_e2e/_fixtures.py
"""

from __future__ import annotations

from pathlib import Path

FIXTURE_DIR = Path(__file__).parent.parent / "fixtures" / "chart_e2e"

# Data：类别 + 两列数值，供柱/折/饼/散点等图表取数。
_DATA_ROWS: list[list[object]] = [
    ["Month", "Sales", "Cost"],
    ["Jan", 100, 60],
    ["Feb", 150, 80],
    ["Mar", 200, 90],
]


def _fill(ws, rows: list[list[object]]) -> None:
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=value)


def build_chart(path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    data = wb.active
    data.title = "Data"
    _fill(data, _DATA_ROWS)
    wb.create_sheet("Blank")  # 空表，供 get:charts 验证空列表
    wb.active = 0  # activeSheet = Data
    wb.save(str(path))


def build_all() -> dict[str, Path]:
    FIXTURE_DIR.mkdir(parents=True, exist_ok=True)
    paths: dict[str, Path] = {}
    for name, builder in {"chart.xlsx": build_chart}.items():
        target = FIXTURE_DIR / name
        builder(target)
        paths[name] = target
        print(f"📝 生成夹具: {target}")
    return paths


def ensure_fixtures() -> dict[str, Path]:
    """确保 chart.xlsx 夹具存在（缺失才生成）。"""
    FIXTURE_DIR.mkdir(parents=True, exist_ok=True)
    if (FIXTURE_DIR / "chart.xlsx").exists():
        return {"chart.xlsx": FIXTURE_DIR / "chart.xlsx"}
    return build_all()


if __name__ == "__main__":
    build_all()
