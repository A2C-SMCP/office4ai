"""table_e2e 夹具构建器。

生成单个多 sheet 夹具 ``tbl.xlsx``（提交入库，供真机打开），覆盖 #33 全部 table 用例：
建表（insert）/ 查表（get:table·get:tables）/ 行增删（add·delete tableRow）/ 排序（sort）。

四张表：

- ``Raw`` —— **activeSheet**，A1:C4 纯数据网格（无 Excel 表）。供 ``insert:table`` 在已有
  数据上建表。
- ``Blank`` —— 空表。供 ``insert:table`` 带 data 写入空白区。
- ``Sales`` —— 预置 Excel 表 ``SalesTable``（A1:C5，表头 Name/Score/Age + 4 行，Score 含并列
  85 供多级排序断 tie）。供 ``get:table`` / ``get:tables`` / ``sort:table``。
- ``Roster`` —— 预置 Excel 表 ``RosterTable``（A1:B4，表头 Item/Qty + 3 行自描述
  R0row/R1row/R2row）。供 ``add:tableRow`` / ``delete:tableRow``（含 rowIndex=0 falsy）。

openpyxl 直接写 Excel 表（ListObject），Excel 打开后 Office.js ``worksheet.tables`` 即可识别。

重新生成：
    uv run python manual_tests/excel/table_e2e/_fixtures.py
"""

from __future__ import annotations

from pathlib import Path

FIXTURE_DIR = Path(__file__).parent.parent / "fixtures" / "table_e2e"

# Raw：纯数据网格（无表）。
_RAW_ROWS: list[list[object]] = [
    ["Region", "Q1", "Q2"],
    ["North", 100, 150],
    ["South", 200, 250],
    ["East", 300, 350],
]

# Sales 表：Score 含并列 85（Bob/Alice）供多级排序（Score↑ 再 Age↑）断 tie；Age 唯一供单列排序。
_SALES_ROWS: list[list[object]] = [
    ["Name", "Score", "Age"],
    ["Bob", 85, 30],
    ["Alice", 85, 25],
    ["Dave", 78, 35],
    ["Carol", 88, 28],
]

# Roster 表：Item 列自描述（R0row/R1row/R2row），便于断言行增删后「谁还在 / 谁上移」。
_ROSTER_ROWS: list[list[object]] = [
    ["Item", "Qty"],
    ["R0row", 10],
    ["R1row", 20],
    ["R2row", 30],
]


def _fill(ws, rows: list[list[object]]) -> None:
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=value)


def _add_table(ws, display_name: str, ref: str) -> None:
    from openpyxl.worksheet.table import Table, TableStyleInfo

    tab = Table(displayName=display_name, ref=ref)
    tab.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium9", showRowStripes=True, showFirstColumn=False, showLastColumn=False
    )
    ws.add_table(tab)


def build_tbl(path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    raw = wb.active
    raw.title = "Raw"
    _fill(raw, _RAW_ROWS)

    wb.create_sheet("Blank")  # 空表，供 insert:table 带 data

    sales = wb.create_sheet("Sales")
    _fill(sales, _SALES_ROWS)
    _add_table(sales, "SalesTable", "A1:C5")

    roster = wb.create_sheet("Roster")
    _fill(roster, _ROSTER_ROWS)
    _add_table(roster, "RosterTable", "A1:B4")

    wb.active = 0  # activeSheet = Raw
    wb.save(str(path))


def build_all() -> dict[str, Path]:
    FIXTURE_DIR.mkdir(parents=True, exist_ok=True)
    paths: dict[str, Path] = {}
    for name, builder in {"tbl.xlsx": build_tbl}.items():
        target = FIXTURE_DIR / name
        builder(target)
        paths[name] = target
        print(f"📝 生成夹具: {target}")
    return paths


def ensure_fixtures() -> dict[str, Path]:
    """确保 tbl.xlsx 夹具存在（缺失才生成）。"""
    FIXTURE_DIR.mkdir(parents=True, exist_ok=True)
    if (FIXTURE_DIR / "tbl.xlsx").exists():
        return {"tbl.xlsx": FIXTURE_DIR / "tbl.xlsx"}
    return build_all()


if __name__ == "__main__":
    build_all()
