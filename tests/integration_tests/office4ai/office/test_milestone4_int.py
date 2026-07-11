"""milestone #4 末端跨边界集成回归 | milestone #4 end-to-end cross-boundary regression (issue #67).

守护「任何单个 sub-task 都无法独立守护」的跨边界不变量——依赖图末端的唯一汇聚点。纯测试、
零生产改动。四类不变量与其覆盖归属：

- ① **W1→W2→W3 链式组装**（本文件 :class:`TestPipelineChainW1W2W3`）：单一 artifact 依次流过
  create（W1 shipped 脚本）→ edit（W2 流水线编辑）→ extract（W3 shipped 脚本），证明 authoring
  流水线跨 SKILL 边界可组合，末端产物保真且可被回填消费。各 SKILL 的独立跑通已由
  ``test_{create,edit,extract}_office_file.py`` 守护——此处专测**组合**。注：**真实模板实例化**
  （``instantiate_from_template`` / ``.potx``）**不在**本组合链，由 S4 ``test_create_office_file.py``
  专测；本链聚焦跨阶段 artifact 连续性，W2 用代表性流水线编辑（非 shipped edit 脚本，理由见其 docstring）。
- ④ **三生产 SKILL 发现**（:class:`TestThreeProductionSkillsDiscovery`）：起真 ``OfficeMCPServer``
  （**默认** skills root，不 override ``OFFICE4AI_SKILLS_ROOT``），断言 ``skill://`` 同时暴露
  create/edit/extract 三个**生产** SKILL 根 + 各自 scripts 子资源，且脚本确实接线到 authoring
  helper 原语库。``test_skill_resource_mcp.py`` 只测 fixture ``demo-authoring``，未覆盖此汇聚。
- ②③ **W4a 工具收敛 + W4b 单 fullscreen**：已由 ``test_w4_desktop_convergence.py``（#63–66 的执行
  手段）端到端守护，本文件不重复、不改动它；仅在 :class:`TestWholeDesktopSnapshot` 里把四类不变量
  **合并进一个 MCP 会话**，守护它们在「Computer 眼中的完整桌面」同时成立（W4×S3 交叉）。

架构立场（路线甲）：office4ai 是 MCP Server（生产者/桌面权威），A2C Computer 是消费者。生产者不得
反向依赖消费者——因此「Computer」一律由**通用 MCP client** 充当消费方观测 MCP wire 契约，绝不 import
python-sdk / A2C Computer（保持 ``python-sdk 零改动 + 单仓``）。
"""

from __future__ import annotations

import os
import zipfile
from pathlib import Path
from unittest.mock import AsyncMock, patch

import pytest
from docx import Document
from mcp.shared.memory import create_connected_server_and_client_session
from mcp.types import TextResourceContents
from pydantic import AnyUrl

import office4ai
from office4ai.a2c_smcp.config import MCPServerConfig
from office4ai.a2c_smcp.resources.skill import DEFAULT_SKILL_HOST
from office4ai.environment.workspace.socketio.services.connection_manager import connection_manager
from office4ai.office.authoring import run_script
from office4ai.office.authoring.helpers import fill_sdt_controls
from office4ai.office.mcp.server import OfficeMCPServer

pytestmark = pytest.mark.integration

SKILLS_ROOT = Path(office4ai.__file__).resolve().parent / "office" / "skills"
CREATE_SCRIPTS = SKILLS_ROOT / "create-office-file" / "scripts"
EXTRACT_SCRIPTS = SKILLS_ROOT / "extract-template" / "scripts"
SKILL_BASE = f"skill://{DEFAULT_SKILL_HOST}"


def _substitute(src: str, replacements: dict[str, str]) -> str:
    """把 shipped 脚本顶部占位常量（``NAME = "..."``）替换成测试真实值——复刻 SKILL 集成测试手法。"""
    import re

    for key, val in replacements.items():
        new_src, n = re.subn(rf'^{key} = ".*"', f"{key} = {val!r}", src, count=1, flags=re.M)
        assert n == 1, f"placeholder constant {key} not found in shipped script"
        src = new_src
    return src


def _docx_xml(path: Path) -> str:
    """读 word/document.xml 原始 XML——SDT 内容嵌在 ``w:sdtContent``，不出现在 ``Document.paragraphs``。"""
    with zipfile.ZipFile(path) as zf:
        return zf.read("word/document.xml").decode("utf-8")


@pytest.fixture
def clean_connection_manager():
    """进程级单例 connection_manager 隔离：清空客户端 + 快照/还原连接回调列表，避免跨用例污染。"""
    for c in list(connection_manager.get_all_clients()):
        connection_manager.unregister_client(c.socket_id)
    saved_connect = list(connection_manager._on_document_connect)
    saved_disc = list(connection_manager._on_document_disconnect_ns)
    try:
        yield
    finally:
        for c in list(connection_manager.get_all_clients()):
            connection_manager.unregister_client(c.socket_id)
        connection_manager._on_document_connect[:] = saved_connect
        connection_manager._on_document_disconnect_ns[:] = saved_disc


# ---------------------------------------------------------------------------
# ① W1→W2→W3 链式组装：单 artifact 流过三阶段（跨 SKILL 边界的组合，非各自孤立）
# ---------------------------------------------------------------------------


class TestPipelineChainW1W2W3:
    async def test_create_edit_extract_pipeline_chain(self, tmp_path: Path) -> None:
        """create 产物 → edit → extract 一条龙；全程 office_run_script 真沙箱，末端产物保真且可回填。

        W1 与 W3 跑 **shipped 参考脚本本体**（捕获 import/签名漂移）；W2 是一段代表性流水线编辑，
        用于给 W1 的纯样式产物注入一个 Heading-1 章节，使 W3 的「按标题样式抽 SDT」有锚点可抽——
        这正是 shipped edit 脚本各自被 ``test_edit_office_file.py`` 覆盖后，INT 独有的**组合**职责。
        """
        chain = tmp_path  # 单一工作目录：artifact 逐阶段在此演进

        # ── W1 create：shipped gen_word.py（从零）→ red_header.docx ──────────────
        w1_src = (CREATE_SCRIPTS / "gen_word.py").read_text(encoding="utf-8")
        r1 = await run_script(w1_src, work_dir=chain)
        assert r1.ok is True, f"W1 create failed: {r1.stderr}"
        assert (chain / "red_header.docx").is_file()

        # ── W2 edit：打开 W1 产物，追加 Heading-1 章节（流水线连续性 + 为 W3 造锚点）──────
        w2_src = (
            "from docx import Document\n"
            "doc = Document('red_header.docx')\n"
            "doc.add_heading('第一章 总体要求', level=1)\n"
            "doc.add_paragraph('本章为编辑阶段新增的正文段落。')\n"
            "doc.save('edited.docx')\n"
            "print('edited -> edited.docx')\n"
        )
        r2 = await run_script(w2_src, work_dir=chain)
        assert r2.ok is True, f"W2 edit failed: {r2.stderr}"
        edited = chain / "edited.docx"
        assert edited.is_file()
        texts2 = [p.text for p in Document(str(edited)).paragraphs]
        assert any("XX 市人民政府办公室" in t for t in texts2), "W1 红头内容须在编辑后保真"
        assert "第一章 总体要求" in texts2, "W2 编辑（新增章节）须落到产物"

        # ── W3 extract：shipped extract_docx_sdt.py（REF=W2 产物）→ 标题段抽成命名 SDT 模板 ──
        # STYLE_ID 显式钉成 "Heading1"（= python-docx add_heading(level=1) 产出的 styleId），把
        # 「shipped 默认恰好匹配」的隐式耦合变成测试内的显式契约——shipped 默认漂移时此处即失配可见。
        w3_src = _substitute(
            (EXTRACT_SCRIPTS / "extract_docx_sdt.py").read_text(encoding="utf-8"),
            {"REF": str(edited), "STYLE_ID": "Heading1"},
        )
        r3 = await run_script(w3_src, work_dir=chain)
        assert r3.ok is True, f"W3 extract failed: {r3.stderr}"
        template = chain / "template.docx"
        assert template.is_file()

        # 抽取产物保真：W2 的标题段被原地升级为命名 SDT 占位（灰显标记 + 占位文本）
        tpl_xml = _docx_xml(template)
        assert "showingPlcHdr" in tpl_xml, "抽出的 SDT 须带占位灰显标记"
        assert "【第一章 总体要求】" in tpl_xml, "占位文本取自被抽标题"

        # 末端闭环（W3 产物 → W2 可消费）：抽出的模板可被 fill_sdt_controls 按 alias 精准回填
        filled = chain / "filled.docx"
        n = fill_sdt_controls(str(template), {"标题1": "新的章节标题"}, output=str(filled))
        assert n == 1, "抽出的模板应恰有 1 个命名 SDT（W2 注入的唯一 Heading-1）可回填"
        filled_xml = _docx_xml(filled)
        assert "新的章节标题" in filled_xml, "回填值须落入 SDT 内容"
        assert "【第一章 总体要求】" not in filled_xml, "占位文本回填后应被替换"


# ---------------------------------------------------------------------------
# ④ 三生产 SKILL 发现：默认 skills root 下 skill:// 暴露 create/edit/extract + helper 接线
# ---------------------------------------------------------------------------


class TestThreeProductionSkillsDiscovery:
    async def test_default_skills_root_exposes_three_production_skills(self) -> None:
        """起真 OfficeMCPServer（默认 skills root），断言三生产 SKILL 根 + scripts + helper 接线均可发现。

        用 MCP 内存 client 充当 A2C Computer 消费方（不 import python-sdk）。断言的是 staging.py 消费的
        wire 契约：根带 ``_meta.source=resources``、子资源可读、脚本接线到 authoring helper 原语库。
        """
        expected_roots = {
            f"{SKILL_BASE}/create-office-file",
            f"{SKILL_BASE}/edit-office-file",
            f"{SKILL_BASE}/extract-template",
        }
        # 默认 skills root：不 override OFFICE4AI_SKILLS_ROOT → 包内 office/skills/
        with patch.dict(os.environ, {}, clear=True):
            office = OfficeMCPServer(MCPServerConfig())

        async with create_connected_server_and_client_session(office.server) as client:
            resources = (await client.list_resources()).resources
            by_uri = {str(r.uri): r for r in resources}

            # 三生产 SKILL 根均暴露，且带 _meta.source=resources（否则消费方不当作 SKILL 根）
            for root_uri in expected_roots:
                assert root_uri in by_uri, f"missing production SKILL root {root_uri}"
                meta = by_uri[root_uri].meta
                assert meta is not None and meta.get("source") == "resources", root_uri

            # 子资源：三 SKILL 的参考脚本作为兄弟资源可发现（Computer 物化脚本包的依据）
            uris = set(by_uri)
            assert f"{SKILL_BASE}/create-office-file/scripts/gen_word.py" in uris
            assert f"{SKILL_BASE}/edit-office-file/scripts/edit_docx_sdt.py" in uris
            assert f"{SKILL_BASE}/extract-template/scripts/extract_docx_sdt.py" in uris

            # + helper：物化的脚本确实接线到 authoring helper 原语库（S2×S3×S6 交叉）。查「行首实义
            # import」而非裸子串——注释掉的 import 不算数（链式 ① 已在沙箱里真实执行该 import 作强证据）。
            got = await client.read_resource(AnyUrl(f"{SKILL_BASE}/extract-template/scripts/extract_docx_sdt.py"))
            content = got.contents[0]
            assert isinstance(content, TextResourceContents)
            assert any(
                line.strip().startswith("from office4ai.office.authoring.helpers import")
                for line in content.text.splitlines()
            ), "shipped 脚本须实际（非注释）接线到 authoring helper 原语库"


# ---------------------------------------------------------------------------
# 整机桌面快照汇聚：一个 MCP 会话内四类不变量并存（W4a 工具收敛 × W4b 窗口/fullscreen × S3 skill://）
# ---------------------------------------------------------------------------


@pytest.mark.usefixtures("clean_connection_manager")
class TestWholeDesktopSnapshot:
    async def test_computer_sees_converged_tools_windows_fullscreen_and_skills(self) -> None:
        """Computer 眼中的完整桌面：三平台连接下，工具收敛 + per-file window×3 + 单 fullscreen + skill://三SKILL 同时成立。

        ②③ 的独立端到端已由 test_w4_desktop_convergence.py 守护；此处专测四者在**同一会话**的汇聚
        （单 sub-task 无法覆盖的交叉），并把 skill:// 暴露叠加进同一张桌面快照。
        """
        with patch.dict(os.environ, {}, clear=True):
            office = OfficeMCPServer(MCPServerConfig())
        # 接线连接/活动回调（生产在 _async_startup 完成；startup 本身的接线由
        # test_async_startup_wires_desktop_callbacks 单独守护，此处聚焦收敛后的桌面聚合）
        connection_manager.register_connect_callback(office._on_doc_connect)
        connection_manager.register_disconnect_callback_ns(office._on_doc_disconnect)
        office.workspace.set_activity_callback(office._on_doc_activity)

        per_file_prefixes = (
            "window://office4ai/word/",
            "window://office4ai/ppt/",
            "window://office4ai/excel/",
        )
        word_doc, ppt_doc, xlsx_doc = "file:///tmp/report.docx", "file:///tmp/deck.pptx", "file:///tmp/data.xlsx"

        async with create_connected_server_and_client_session(office.server) as client:
            # 收敛基线（负向）：无连接 → 仅常驻工具、无 per-file 窗口。有此基线，下方「三平台工具全在」
            # 才真正证明「按连接收敛」（空→满），而非「工具本就恒在」。
            base_tools = {t.name for t in (await client.list_tools()).tools}
            assert base_tools == {"office_run_script"}, f"无连接基线应仅常驻工具，实得 {base_tools}"
            base_uris = {str(r.uri) for r in (await client.list_resources()).resources}
            assert not any(u.startswith(per_file_prefixes) for u in base_uris), "无连接不应有 per-file 窗口"

            connection_manager.register_client("s1", "c1", word_doc, "/word")
            connection_manager.register_client("s2", "c2", ppt_doc, "/ppt")
            connection_manager.register_client("s3", "c3", xlsx_doc, "/excel")
            # AI 最后操作 ppt 文件 → fullscreen 应独占归属 ppt 窗口
            office.workspace.update_last_activity(ppt_doc, "ppt_insert_text", {})

            tools = {t.name for t in (await client.list_tools()).tools}
            resources = (await client.list_resources()).resources
            uris = {str(r.uri) for r in resources}

            # ① 工具收敛：常驻 office_run_script + 三平台 Add-In 工具全暴露（相较上方空基线，证明收敛方向）
            assert "office_run_script" in tools
            assert any(n.startswith("word_") for n in tools), "Word 连接 → Word 工具应暴露"
            assert any(n.startswith("ppt_") for n in tools), "PPT 连接 → PPT 工具应暴露"
            assert any(n.startswith("excel_") for n in tools), "Excel 连接 → Excel 工具应暴露"

            # ② per-file window ×3（word/ppt/excel 各一）
            assert any(u.startswith("window://office4ai/word/") for u in uris)
            assert any(u.startswith("window://office4ai/ppt/") for u in uris)
            assert any(u.startswith("window://office4ai/excel/") for u in uris)

            # ③ 单 fullscreen：恰一个 per-file 窗口 fullscreen，且归属 AI 最后操作的 ppt（#4 不变量）
            fs = {u for u in uris if u.startswith(per_file_prefixes) and "fullscreen=true" in u}
            assert len(fs) == 1, f"至多且恰一个 fullscreen 窗口，实得 {fs}"
            assert any("/ppt/" in u for u in fs), "fullscreen 应归属最后被操作的 ppt"

            # ④ skill:// 三生产 SKILL 与桌面窗口同时可见（S3 × W4 交叉）
            for name in ("create-office-file", "edit-office-file", "extract-template"):
                assert f"{SKILL_BASE}/{name}" in uris, f"生产 SKILL {name} 应与窗口并存于同一资源快照"

    async def test_async_startup_wires_desktop_callbacks(self) -> None:
        """守护启动接线本身（#67 汇聚点补强）：真实 ``_async_startup`` 必须把三个桌面回调接上。

        快照测试手工复刻接线，故无法捕获「_async_startup 漏接某回调」的回归。此处驱动真实
        ``_async_startup``（仅 mock 掉会绑端口的 ``workspace.start``），断言连接/断连回调进了全局
        connection_manager、活动回调（fullscreen 归属触发源）接到 ``_on_doc_activity``。
        """
        with patch.dict(os.environ, {}, clear=True):
            office = OfficeMCPServer(MCPServerConfig())

        with (
            patch.object(office.workspace, "start", new=AsyncMock()),
            patch.object(office.workspace, "set_activity_callback") as set_activity,
        ):
            await office._async_startup()

        assert office._on_doc_connect in connection_manager._on_document_connect, "startup 须注册连接回调"
        assert office._on_doc_disconnect in connection_manager._on_document_disconnect_ns, "startup 须注册断连回调"
        set_activity.assert_any_call(office._on_doc_activity)  # fullscreen 归属触发源必须接线
