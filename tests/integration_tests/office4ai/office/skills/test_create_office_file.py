"""create-office-file SKILL 集成回归 | create-office-file SKILL integration test (milestone #4 · S4).

执行手段（Issue #60 验收）：把 SKILL 的参考脚本经 **office_run_script**（真沙箱子进程）跑通，
断言三平台成品有效；并验证模板复用链路（.potx content-type 互换，无需 soffice）。

真依赖：真实子进程 + 真实 python-docx/pptx/openpyxl + 真实文件系统；不 mock 运行时。
"""

from __future__ import annotations

from pathlib import Path

import docx
import openpyxl
import pytest
from pptx import Presentation

import office4ai
from office4ai.office.authoring import run_script
from office4ai.office.authoring.helpers._ooxml import CT_POTX, CT_PPTX, read_package, write_package

pytestmark = pytest.mark.integration

SKILL_DIR = Path(office4ai.__file__).resolve().parent / "office" / "skills" / "create-office-file"
SCRIPTS = SKILL_DIR / "scripts"


async def _run_ref_script(script_name: str, work_dir: Path) -> Path:
    """把 scripts/<name> 经 office_run_script 沙箱跑通，返回产物路径（断言 ok）。"""
    src = (SCRIPTS / script_name).read_text(encoding="utf-8")
    result = await run_script(src, work_dir=work_dir)
    assert result.ok is True, f"{script_name} failed: {result.stderr}"
    produced = result.summary["produced"]
    assert produced, f"{script_name} produced nothing: {result.logs} / {result.stderr}"
    return work_dir / produced[0]["name"]


async def test_gen_word_from_scratch(tmp_path: Path) -> None:
    out = await _run_ref_script("gen_word.py", tmp_path)
    assert out.name == "red_header.docx"
    texts = [p.text for p in docx.Document(str(out)).paragraphs]
    assert any("XX 市人民政府办公室" in t for t in texts)
    assert any("关于开展 2026 年度工作的通知" in t for t in texts)


async def test_gen_ppt_from_scratch(tmp_path: Path) -> None:
    out = await _run_ref_script("gen_ppt.py", tmp_path)
    assert out.name == "sales_report.pptx"
    prs = Presentation(str(out))
    assert len(prs.slides) == 1
    shapes = prs.slides[0].shapes
    assert any(s.has_chart for s in shapes), "expected a chart on the slide"
    assert any(s.has_table for s in shapes), "expected a table on the slide"


async def test_gen_excel_from_scratch(tmp_path: Path) -> None:
    out = await _run_ref_script("gen_excel.py", tmp_path)
    assert out.name == "sales_book.xlsx"
    wb = openpyxl.load_workbook(str(out))
    ws = wb["销售数据"]
    assert ws["A1"].value == "月份"
    assert ws["E2"].value == "=SUM(B2:B5)-SUM(C2:C5)"
    assert "销售额区" in wb.defined_names


async def test_template_reuse_potx_no_soffice(tmp_path: Path) -> None:
    # 程序化造一份自制 .potx（真实 pptx 经 presentation→template content-type 互换，纯 zip、无 soffice）
    src_pptx = tmp_path / "src.pptx"
    Presentation().save(str(src_pptx))
    infos, items = read_package(src_pptx)
    items["[Content_Types].xml"] = items["[Content_Types].xml"].replace(CT_PPTX.encode(), CT_POTX.encode())
    template = tmp_path / "brand.potx"
    write_package(template, infos, items)

    work = tmp_path / "wd"
    work.mkdir()
    script = (
        "from office4ai.office.authoring.helpers import instantiate_from_template\n"
        f"instantiate_from_template({str(template)!r}, 'deck.pptx')\n"
        "print('instantiated deck.pptx')\n"
    )
    result = await run_script(script, work_dir=work, template_uri=str(template))

    assert result.ok is True, result.stderr
    deck = work / "deck.pptx"
    assert deck.exists()
    Presentation(str(deck))  # 可被 python-pptx 正常打开 == 有效 pptx


async def test_template_reuse_instantiate_and_fill(tmp_path: Path) -> None:
    """头牌链路端到端（沙箱内）：instantiate 模板 + 用 helper 填充（新建自定义版式）。CI 友好（potx，无 soffice）。"""
    src_pptx = tmp_path / "src.pptx"
    Presentation().save(str(src_pptx))
    infos, items = read_package(src_pptx)
    items["[Content_Types].xml"] = items["[Content_Types].xml"].replace(CT_PPTX.encode(), CT_POTX.encode())
    template = tmp_path / "brand.potx"
    write_package(template, infos, items)

    work = tmp_path / "wd"
    work.mkdir()
    # 脚本在沙箱里：instantiate_from_template → 打开成品 → create_custom_layout 填充 → save
    script = (
        "from pptx import Presentation\n"
        "from office4ai.office.authoring.helpers import create_custom_layout, instantiate_from_template\n"
        f"instantiate_from_template({str(template)!r}, 'deck.pptx')\n"
        "prs = Presentation('deck.pptx')\n"
        "create_custom_layout(prs, '品牌版式', lambda shapes: None)\n"
        "prs.save('deck.pptx')\n"
        "print('instantiated + filled deck.pptx')\n"
    )
    result = await run_script(script, work_dir=work, template_uri=str(template))

    assert result.ok is True, result.stderr
    prs = Presentation(str(work / "deck.pptx"))
    layout_names = [layout.name for master in prs.slide_masters for layout in master.slide_layouts]
    assert "品牌版式" in layout_names, f"custom layout not persisted: {layout_names}"
