"""edit-office-file SKILL 集成回归 | edit-office-file SKILL integration test (milestone #4 · S5).

执行手段（Issue #61 验收）：以**程序化生成**的模板/业务文件为素材（不提交二进制），把 SKILL **shipped
参考脚本本体**（`scripts/*.py`，仅替换其占位路径常量）经 office_run_script 真沙箱跑通，断言**模板操作
保真**：docxtpl `{{}}` 渲染 / SDT 填充 / 保图表改数据 / 母版复用。跑 shipped 脚本本体（而非平行内联副本）
可捕获参考脚本的 import/签名漂移。真依赖：真子进程 + 真 docx/pptx/openpyxl/docxtpl/lxml。
"""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

import openpyxl
import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl.chart import BarChart, Reference
from pptx import Presentation

import office4ai
from office4ai.office.authoring import ScriptResult, run_script

pytestmark = pytest.mark.integration

SCRIPTS = Path(office4ai.__file__).resolve().parent / "office" / "skills" / "edit-office-file" / "scripts"


def _substitute(src: str, replacements: dict[str, str]) -> str:
    """把 shipped 脚本顶部的占位常量（TARGET/TEMPLATE/LAYOUT_NAME = "..."）替换成测试真实值。"""
    for key, val in replacements.items():
        new_src, n = re.subn(rf'^{key} = ".*"', f"{key} = {val!r}", src, count=1, flags=re.M)
        assert n == 1, f"placeholder constant {key} not found in shipped script"
        src = new_src
    return src


async def _run_shipped(name: str, work_dir: Path, replacements: dict[str, str], template_uri: str) -> ScriptResult:
    src = _substitute((SCRIPTS / name).read_text(encoding="utf-8"), replacements)
    result = await run_script(src, work_dir=work_dir, template_uri=template_uri)
    assert result.ok is True, f"{name} failed: {result.stderr}"
    return result


def _wrap_in_sdt(paragraph, alias: str, sdt_id: int) -> None:
    """把段落升级成 SDT 内容控件（含 showingPlcHdr 占位标记）——复刻 S2 夹具构造手法。"""
    p = paragraph._p
    sdt = OxmlElement("w:sdt")
    sdt_pr = OxmlElement("w:sdtPr")
    for tag, val in (("w:alias", alias), ("w:tag", alias), ("w:id", str(sdt_id))):
        el = OxmlElement(tag)
        el.set(qn("w:val"), val)
        sdt_pr.append(el)
    sdt_pr.append(OxmlElement("w:showingPlcHdr"))  # 占位灰显标记，填充后应被 helper 剥离
    sdt.append(sdt_pr)
    content = OxmlElement("w:sdtContent")
    parent = p.getparent()
    idx = parent.index(p)
    parent.remove(p)
    content.append(p)
    sdt.append(content)
    parent.insert(idx, sdt)


def _docx_xml(path: Path) -> str:
    with zipfile.ZipFile(path) as zf:
        return zf.read("word/document.xml").decode("utf-8")


async def test_edit_docx_docxtpl_render(tmp_path: Path) -> None:
    tpl = tmp_path / "letter_template.docx"
    doc = Document()
    doc.add_paragraph("尊敬的 {{name}}，本次金额 {{amount}} 元。")
    doc.save(str(tpl))

    work = tmp_path / "wd"
    work.mkdir()
    result = await _run_shipped("edit_docx_docxtpl.py", work, {"TEMPLATE": str(tpl)}, str(tpl))

    out = work / result.summary["produced"][0]["name"]
    texts = " ".join(p.text for p in Document(str(out)).paragraphs)
    assert "张三" in texts and "12,800" in texts
    assert "{{" not in texts  # 占位已替换，无残留


async def test_edit_docx_sdt_fill(tmp_path: Path) -> None:
    # shipped 脚本默认 strict=True 填 3 个 alias → 夹具须含全部 3 个（文档标题/副标题/章节一标题）
    src = tmp_path / "report_with_sdt.docx"
    doc = Document()
    for i, (text, alias) in enumerate(
        [("【文档标题】", "文档标题"), ("【副标题】", "副标题"), ("【章节一标题】", "章节一标题")],
    ):
        para = doc.add_paragraph()
        para.add_run(text)
        _wrap_in_sdt(para, alias=alias, sdt_id=1001 + i)
    doc.save(str(src))

    work = tmp_path / "wd"
    work.mkdir()
    result = await _run_shipped("edit_docx_sdt.py", work, {"TARGET": str(src)}, str(src))

    xml = _docx_xml(work / result.summary["produced"][0]["name"])
    assert "2026 年度工作报告" in xml  # SDT 内容嵌在 w:sdtContent，读原始 XML
    assert "【文档标题】" not in xml  # 占位文本被替换
    assert "showingPlcHdr" not in xml  # 填充后占位灰显标记被剥离（保真行为）


async def test_edit_xlsx_preserve_chart(tmp_path: Path) -> None:
    src = tmp_path / "dashboard.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "数据"
    ws["A1"], ws["B1"] = "项目", "金额"
    ws["A2"], ws["B2"] = "收入", 100
    ws["A3"], ws["B3"] = "支出", 60
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=3))
    ws.add_chart(chart, "D1")
    wb.save(str(src))

    work = tmp_path / "wd"
    work.mkdir()
    result = await _run_shipped("edit_xlsx_preserve_chart.py", work, {"TEMPLATE": str(src)}, str(src))

    out = work / result.summary["produced"][0]["name"]
    with zipfile.ZipFile(out) as zf:  # 保真：图表 XML 仍在包内（fill_cells_lxml 外科式编辑保全）
        assert any(n.startswith("xl/charts/") for n in zf.namelist()), "chart lost after edit"
    assert openpyxl.load_workbook(out)["数据"]["B2"].value == 260  # 数据已改
    # 注：不对 openpyxl 整体重存丢图表做断言——该风险与 openpyxl 版本/图表复杂度相关（简单图表可能幸存），
    # SKILL 文案已据此措辞为「不完整往返有丢失风险」而非绝对「会丢」。


async def test_edit_ppt_master_reuse(tmp_path: Path) -> None:
    base = Presentation()
    layout_name = base.slide_masters[0].slide_layouts[1].name  # 默认模板里真实版式名（英文）
    src = tmp_path / "branded_deck.pptx"
    base.save(str(src))

    work = tmp_path / "wd"
    work.mkdir()
    result = await _run_shipped(
        "edit_ppt_master.py",
        work,
        {"TARGET": str(src), "LAYOUT_NAME": layout_name},
        str(src),
    )

    out = Presentation(str(work / result.summary["produced"][0]["name"]))
    assert len(out.slides) == 1
    assert out.slides[0].slide_layout.name == layout_name  # 新页复用了模板母版版式
