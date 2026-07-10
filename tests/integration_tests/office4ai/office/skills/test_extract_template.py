"""extract-template SKILL 集成回归 | extract-template SKILL integration test (milestone #4 · S6).

执行手段（Issue #62 验收）：以**两个真实 .docx** 为参考文件（程序化生成，不提交二进制），把 SKILL **shipped
参考脚本本体**（`scripts/*.py`，仅替换其占位路径常量）经 office_run_script 真沙箱跑通，断言**参考文件→可复用
模板**且产物**能被 W1/W2 消费**——抽取↔填充闭环：`wrap_in_sdt`→`fill_sdt_controls`、`{{token}}`→docxtpl 渲染。
另附 PPT 版式抽取 + Excel 命名区域抽取的保真冒烟。跑 shipped 脚本本体（而非平行内联副本）可捕获参考脚本的
import/签名漂移。真依赖：真子进程 + 真 docx/pptx/openpyxl/docxtpl/lxml。
"""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

import openpyxl
import pytest
from docx import Document
from docxtpl import DocxTemplate
from openpyxl.chart import BarChart, Reference
from openpyxl.workbook.defined_name import DefinedName
from pptx import Presentation
from pptx.util import Emu

import office4ai
from office4ai.office.authoring import ScriptResult, run_script
from office4ai.office.authoring.helpers import fill_sdt_controls, locate_by_master, locate_by_named_range
from office4ai.office.authoring.helpers._ooxml import parse_xml, qn, read_package

pytestmark = pytest.mark.integration

SCRIPTS = Path(office4ai.__file__).resolve().parent / "office" / "skills" / "extract-template" / "scripts"


def _substitute(src: str, replacements: dict[str, str]) -> str:
    """把 shipped 脚本顶部的占位常量（REF = "..."）替换成测试真实值。"""
    for key, val in replacements.items():
        new_src, n = re.subn(rf'^{key} = ".*"', f"{key} = {val!r}", src, count=1, flags=re.M)
        assert n == 1, f"placeholder constant {key} not found in shipped script"
        src = new_src
    return src


async def _run_shipped(name: str, work_dir: Path, replacements: dict[str, str], template_uri: str) -> ScriptResult:
    src = _substitute((SCRIPTS / name).read_text(encoding="utf-8"), replacements)
    result = await run_script(src, work_dir=work_dir, template_uri=template_uri)
    assert result.ok is True, f"{name} failed: {result.stderr}"
    return result


def _sdt_texts(docx: Path) -> list[str]:
    _, items = read_package(docx)
    doc = parse_xml(items["word/document.xml"])
    out: list[str] = []
    for sdt in doc.iter(qn("w:sdt")):
        content = sdt.find(qn("w:sdtContent"))
        if content is None:
            continue
        texts = list(content.iter(qn("w:t")))
        if texts:
            out.append("".join(t.text or "" for t in texts))
    return out


# ── 真实参考 .docx #1：含 Heading1 标题段（抽成 SDT 命名占位）─────────────────
async def test_extract_docx_sdt_then_fill_roundtrip(tmp_path: Path) -> None:
    ref = tmp_path / "reference_report.docx"
    doc = Document()
    for title in ("第一章 概述", "第二章 方法"):
        para = doc.add_paragraph()
        para.style = "Heading 1"
        para.add_run(title)
    doc.add_paragraph("正文段落，非标题。")
    doc.save(str(ref))

    work = tmp_path / "wd"
    work.mkdir()
    await _run_shipped("extract_docx_sdt.py", work, {"REF": str(ref)}, str(ref))

    tpl = work / "template.docx"
    assert tpl.is_file()
    # 抽取产物：两个标题段被升级为 SDT 命名占位（灰显）
    assert _sdt_texts(tpl) == ["【第一章 概述】", "【第二章 方法】"]
    with zipfile.ZipFile(tpl) as zf:
        assert "showingPlcHdr" in zf.read("word/document.xml").decode("utf-8")

    # W2 消费：产出的模板可被 fill_sdt_controls 按 alias 回填（抽取↔填充闭环）
    filled = work / "filled.docx"
    n = fill_sdt_controls(tpl, {"标题1": "新章节甲", "标题2": "新章节乙"}, output=filled)
    assert n == 2
    assert _sdt_texts(filled) == ["新章节甲", "新章节乙"]


# ── 真实参考 .docx #2：含具体实值、且实值被拆进多个 run（抽成 {{token}}）──────
async def test_extract_docx_token_then_docxtpl_roundtrip(tmp_path: Path) -> None:
    ref = tmp_path / "reference_letter.docx"
    doc = Document()
    para = doc.add_paragraph()
    # 实值「张三」「2026-01-01」「12,800」故意拆进多个 run —— 逼真的 run-splitting
    for chunk in ("尊敬的 ", "张", "三", " 先生，签署于 ", "2026", "-01-01", "，金额 ", "12", ",800", " 元。"):
        para.add_run(chunk)
    doc.save(str(ref))

    work = tmp_path / "wd"
    work.mkdir()
    await _run_shipped("extract_docx_token.py", work, {"REF": str(ref)}, str(ref))

    tpl = work / "template.docx"
    assert tpl.is_file()
    body = " ".join(p.text for p in Document(str(tpl)).paragraphs)
    # 抽取产物：实值已换成 docxtpl 占位，且 token 未被拆断（跨 run 合并）
    assert "{{name}}" in body and "{{date}}" in body and "{{amount}}" in body
    assert "张三" not in body

    # W2 消费：docxtpl 渲染回填（证明 token 落在单个 run、可被 Jinja 解析）
    rendered = work / "rendered.docx"
    tpl_doc = DocxTemplate(str(tpl))
    tpl_doc.render({"name": "李四", "date": "2026-07-09", "amount": "9,999"})
    tpl_doc.save(str(rendered))
    out = " ".join(p.text for p in Document(str(rendered)).paragraphs)
    assert "李四" in out and "9,999" in out
    assert "{{" not in out  # 占位已渲染，无残留


# ── PPT：样板页抽成可复用母版版式（保真冒烟）────────────────────────────────
async def test_extract_ppt_layout_reusable(tmp_path: Path) -> None:
    ref = tmp_path / "reference_deck.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Emu(914400), Emu(914400), Emu(3000000), Emu(500000))
    box.text_frame.paragraphs[0].add_run().text = "SAMPLE_MARK"
    prs.save(str(ref))

    work = tmp_path / "wd"
    work.mkdir()
    await _run_shipped("extract_ppt_layout.py", work, {"REF": str(ref)}, str(ref))

    tpl = work / "template.pptx"
    assert tpl.is_file()
    reloaded = Presentation(str(tpl))
    layout = locate_by_master(reloaded, "品牌版式")  # 抽出的版式在
    text = "".join(
        run.text
        for shape in layout.shapes
        if shape.has_text_frame
        for para in shape.text_frame.paragraphs
        for run in para.runs
    )
    assert "SAMPLE_MARK" in text
    # W1/W2 消费：用抽出的版式加页，deck 可用
    reloaded.slides.add_slide(layout)
    used = work / "used.pptx"
    reloaded.save(str(used))
    assert len(Presentation(str(used)).slides) == 1


# ── Excel：命名区域抽成「留结构、清实值」模板（保图表冒烟）────────────────────
async def test_extract_xlsx_named_preserves_chart(tmp_path: Path) -> None:
    ref = tmp_path / "reference_dashboard.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "数据"
    ws["A1"], ws["B1"] = "项目", "金额"
    ws["A2"], ws["B2"] = "收入", 100
    ws["A3"], ws["B3"] = "支出", 60
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=3))
    ws.add_chart(chart, "D1")
    wb.defined_names["金额区"] = DefinedName("金额区", attr_text="数据!$B$2:$B$3")
    wb.save(str(ref))

    work = tmp_path / "wd"
    work.mkdir()
    await _run_shipped("extract_xlsx_named.py", work, {"REF": str(ref)}, str(ref))

    tpl = work / "template.xlsx"
    assert tpl.is_file()
    with zipfile.ZipFile(tpl) as zf:  # 保真：图表 XML 仍在（fill_cells_lxml 外科式清值）
        assert any(nm.startswith("xl/charts/") for nm in zf.namelist()), "chart lost during extraction"
    reloaded = openpyxl.load_workbook(tpl)
    assert reloaded["数据"]["B2"].value is None and reloaded["数据"]["B3"].value is None  # 示例值已清空
    assert locate_by_named_range(tpl, "金额区")  # 命名区域结构仍在，可被回填复用
