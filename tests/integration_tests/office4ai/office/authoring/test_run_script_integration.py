"""authoring 运行时集成回归 | authoring runtime integration test (milestone #4 · S1).

执行手段（Issue #57 验收）：**无 Add-In 连接**下，通过 office_run_script 工具提交一段
python-docx 建 docx 的脚本，跑通并回传结构化结果；同时端到端验证死循环 / 越权写被沙箱拦截。

真依赖：真实子进程 + 真实 python-docx + 真实文件系统（不 mock 运行时）。workspace 依赖
本工具不使用，故以 MagicMock 满足构造即可。
"""

from __future__ import annotations

from pathlib import Path
from unittest.mock import MagicMock

import pytest

from office4ai.a2c_smcp.tools.authoring import OfficeRunScriptTool

pytestmark = pytest.mark.integration


def _tool() -> OfficeRunScriptTool:
    return OfficeRunScriptTool(MagicMock())


async def test_create_docx_without_addin_connection(tmp_path: Path) -> None:
    script = (
        "import docx\n"
        "doc = docx.Document()\n"
        "doc.add_heading('Office4AI Authoring Runtime', level=1)\n"
        "doc.add_paragraph('Generated in the sandbox with no Add-In connection.')\n"
        "doc.save('report.docx')\n"
    )
    result = await _tool().execute({"script": script, "work_dir": str(tmp_path)})

    # 协议验证：结构化返回契约
    assert result["ok"] is True, result["stderr"]
    assert Path(result["path"]) == tmp_path
    produced = [item["name"] for item in result["summary"]["produced"]]
    assert "report.docx" in produced

    # 内容验证：重开产物断言真实结构
    import docx

    reopened = docx.Document(str(tmp_path / "report.docx"))
    texts = [p.text for p in reopened.paragraphs]
    assert "Office4AI Authoring Runtime" in texts
    assert any("no Add-In connection" in t for t in texts)


async def test_infinite_loop_intercepted_through_tool(tmp_path: Path) -> None:
    result = await _tool().execute(
        {"script": "while True:\n    pass\n", "work_dir": str(tmp_path), "timeout": 1.0},
    )
    assert result["ok"] is False
    assert "timeout" in result["stderr"].lower()


async def test_out_of_bounds_write_intercepted_through_tool(tmp_path: Path) -> None:
    outside = tmp_path.parent / "escaped_by_authoring_runtime.txt"
    result = await _tool().execute(
        {"script": f"open({str(outside)!r}, 'w').write('x')\n", "work_dir": str(tmp_path)},
    )
    assert result["ok"] is False
    assert "[sandbox]" in result["stderr"]
    assert not outside.exists()
