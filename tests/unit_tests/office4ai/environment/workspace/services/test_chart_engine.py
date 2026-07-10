"""Unit tests for the OOXML chart engine (in-memory base64 / Path B).

The on-disk "Path A" surface (``insert_chart`` / ``get_chart`` / ``update_chart``
taking a ``document_uri``) was removed in F1 (#68). These tests exercise the sole
remaining engine surface: the in-memory base64 single-slide functions
(``generate_chart_slide_base64`` / ``insert_chart_into_slide_base64`` /
``get_chart_from_slide_base64`` / ``update_chart_in_slide_base64``).
"""

from __future__ import annotations

import base64
import io
import math
from typing import Any

import pytest
from pptx import Presentation

from office4ai.environment.workspace.dtos.common import ErrorCode
from office4ai.environment.workspace.dtos.ppt import (
    CategoricalChartData,
    CategoricalChartUpdate,
    CategoricalSeries,
    ChartInsertOptions,
    ScatterChartData,
    ScatterChartUpdate,
    ScatterPoint,
    ScatterSeries,
)
from office4ai.environment.workspace.services import chart_engine
from office4ai.environment.workspace.services.chart_engine import ChartEngineError

# NOTE: Two Path-A-only error paths were dropped when Path A was removed in #68:
# 3001 DOCUMENT_NOT_FOUND (unknown document URI) and 4002 INVALID_PARAM (on-disk
# multi-slide index out of range). Path A 已在 #68 删除，file-not-found /
# on-disk 多页索引校验随之移除 — the base64 engine has no on-disk file to miss and
# operates on a single slide, so neither check has a base64 equivalent.


# ---------------------------------------------------------------------------
# generate_chart_slide_base64 (new-page generation; was insert_chart Path A)
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_insert_uses_explicit_geometry() -> None:
    chart = CategoricalChartData(
        chartType="Pie",
        categories=["A", "B"],
        series=[CategoricalSeries(name="s", values=[1, 2])],
    )
    result = await chart_engine.generate_chart_slide_base64(
        chart,
        ChartInsertOptions(left=72, top=36, width=400, height=300),
    )
    assert math.isclose(result["left"], 72, abs_tol=0.5)
    assert math.isclose(result["top"], 36, abs_tol=0.5)
    assert math.isclose(result["width"], 400, abs_tol=0.5)
    assert math.isclose(result["height"], 300, abs_tol=0.5)
    assert result["slideIndex"] == 0  # base64 packages are single-page


@pytest.mark.asyncio
async def test_insert_categorical_dimension_mismatch_raises_3015() -> None:
    chart = CategoricalChartData(
        chartType="Pie",
        categories=["A", "B", "C"],
        series=[CategoricalSeries(name="x", values=[1, 2])],  # 2 != 3
    )
    with pytest.raises(ChartEngineError) as exc:
        await chart_engine.generate_chart_slide_base64(chart, None)
    assert exc.value.code == ErrorCode.INVALID_CHART_DATA


@pytest.mark.asyncio
async def test_insert_scatter_non_finite_raises_3015() -> None:
    chart = ScatterChartData(
        chartType="Scatter",
        series=[ScatterSeries(name="s", points=[ScatterPoint(x=float("inf"), y=1)])],
    )
    with pytest.raises(ChartEngineError) as exc:
        await chart_engine.generate_chart_slide_base64(chart, None)
    assert exc.value.code == ErrorCode.INVALID_CHART_DATA


# ---------------------------------------------------------------------------
# get_chart_from_slide_base64 (was get_chart Path A)
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_get_categorical_round_trip() -> None:
    chart = CategoricalChartData(
        chartType="ColumnClustered",
        categories=["Q1", "Q2"],
        series=[
            CategoricalSeries(name="Rev", values=[100, 200]),
            CategoricalSeries(name="Cost", values=[80, 90]),
        ],
        title="Year",
        showLegend=True,
    )
    gen = await chart_engine.generate_chart_slide_base64(chart)
    eid = gen["elementId"]

    got = await chart_engine.get_chart_from_slide_base64(gen["slideBase64"], eid)
    assert got["elementId"] == eid
    assert got["chart"]["chartType"] == "ColumnClustered"
    assert got["chart"]["categories"] == ["Q1", "Q2"]
    assert got["chart"]["title"] == "Year"
    assert len(got["chart"]["series"]) == 2
    assert got["chart"]["series"][0]["name"] == "Rev"
    assert got["chart"]["series"][0]["values"] == [100.0, 200.0]


@pytest.mark.asyncio
async def test_get_unknown_element_raises_3010() -> None:
    with pytest.raises(ChartEngineError) as exc:
        await chart_engine.get_chart_from_slide_base64(_blank_slide_base64(), "chart-99999")
    assert exc.value.code == ErrorCode.ELEMENT_NOT_FOUND


@pytest.mark.asyncio
async def test_get_invalid_element_id_format_raises_3010() -> None:
    with pytest.raises(ChartEngineError) as exc:
        await chart_engine.get_chart_from_slide_base64(_blank_slide_base64(), "shape-1")
    assert exc.value.code == ErrorCode.ELEMENT_NOT_FOUND


# ---------------------------------------------------------------------------
# update_chart_in_slide_base64 (was update_chart Path A)
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_update_categorical_dimension_mismatch_raises_3015() -> None:
    chart = CategoricalChartData(
        chartType="ColumnClustered",
        categories=["A", "B"],
        series=[CategoricalSeries(name="s", values=[1, 2])],
    )
    gen = await chart_engine.generate_chart_slide_base64(chart)

    upd = CategoricalChartUpdate(
        chartType="ColumnClustered",
        categories=["X", "Y", "Z"],
        series=[CategoricalSeries(name="t", values=[10, 20])],  # 2 != 3
    )
    with pytest.raises(ChartEngineError) as exc:
        await chart_engine.update_chart_in_slide_base64(gen["slideBase64"], gen["elementId"], upd)
    assert exc.value.code == ErrorCode.INVALID_CHART_DATA


@pytest.mark.asyncio
async def test_cross_variant_scatter_to_categorical_requires_categories_and_series() -> None:
    sc = ScatterChartData(
        chartType="Scatter",
        series=[ScatterSeries(name="ads", points=[ScatterPoint(x=1, y=2)])],
    )
    gen = await chart_engine.generate_chart_slide_base64(sc)

    # Only chartType — no series, no categories — should fail with 3015.
    upd = CategoricalChartUpdate(chartType="ColumnClustered")
    with pytest.raises(ChartEngineError) as exc:
        await chart_engine.update_chart_in_slide_base64(gen["slideBase64"], gen["elementId"], upd)
    assert exc.value.code == ErrorCode.INVALID_CHART_DATA


@pytest.mark.asyncio
async def test_update_unknown_element_raises_3010() -> None:
    upd = CategoricalChartUpdate(chartType="ColumnClustered", title="x")
    with pytest.raises(ChartEngineError) as exc:
        await chart_engine.update_chart_in_slide_base64(_blank_slide_base64(), "chart-99999", upd)
    assert exc.value.code == ErrorCode.ELEMENT_NOT_FOUND


# ---------------------------------------------------------------------------
# elementId opacity (oasp-chart-<uuid> written to cNvPr/@name) — #13
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_insert_element_id_is_opaque_and_written_to_shape_name() -> None:
    """Generated chart's elementId is opaque and persisted as the shape name (cNvPr/@name)."""
    chart = CategoricalChartData(
        chartType="Line",
        categories=["A", "B"],
        series=[CategoricalSeries(name="s", values=[1, 2])],
    )
    gen = await chart_engine.generate_chart_slide_base64(chart)
    eid = gen["elementId"]
    assert eid.startswith("oasp-chart-")
    # The opaque id must survive encode/decode as the shape's OOXML name.
    prs = Presentation(io.BytesIO(base64.b64decode(gen["slideBase64"])))
    names = [shape.name for slide in prs.slides for shape in slide.shapes if getattr(shape, "has_chart", False)]
    assert eid in names


@pytest.mark.asyncio
async def test_get_relocates_by_opaque_name_not_native_id() -> None:
    """After a second chart shifts native ids, get still finds the first chart by its opaque name."""
    chart = CategoricalChartData(
        chartType="ColumnClustered",
        categories=["A", "B"],
        series=[CategoricalSeries(name="s", values=[1, 2])],
        title="first",
    )
    gen = await chart_engine.generate_chart_slide_base64(chart)
    first = gen["elementId"]
    # Insert a second chart onto the same slide (native shape ids advance).
    second = await chart_engine.insert_chart_into_slide_base64(gen["slideBase64"], chart)

    got = await chart_engine.get_chart_from_slide_base64(second["slideBase64"], first)
    assert got["elementId"] == first
    assert got["chart"]["title"] == "first"


@pytest.mark.asyncio
async def test_legacy_chart_element_id_still_resolves() -> None:
    """A 0.2.0 chart (default python-pptx name, no opaque id) resolves via legacy chart-<slide>-<shape>."""
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Pt

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    cd = CategoryChartData()
    cd.categories = ["A", "B"]
    cd.add_series("s", (1, 2))
    shape: Any = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Pt(10), Pt(10), Pt(300), Pt(200), cd)  # type: ignore[arg-type]
    legacy_eid = f"chart-0-{int(shape.shape_id)}"  # not written to shape.name
    buf = io.BytesIO()
    prs.save(buf)
    slide_b64 = base64.b64encode(buf.getvalue()).decode("ascii")

    got = await chart_engine.get_chart_from_slide_base64(slide_b64, legacy_eid)
    assert got["chart"]["chartType"] == "ColumnClustered"
    assert got["chart"]["categories"] == ["A", "B"]


@pytest.mark.asyncio
async def test_legacy_chart_migrates_to_opaque_id_on_recreate() -> None:
    """A 0.2.0 (legacy-named) chart updated via a cross-variant recreate is migrated to a fresh opaque id."""
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Pt

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    cd = CategoryChartData()
    cd.categories = ["A", "B"]
    cd.add_series("s", (1, 2))
    shape: Any = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Pt(10), Pt(10), Pt(300), Pt(200), cd)  # type: ignore[arg-type]
    legacy_eid = f"chart-0-{int(shape.shape_id)}"  # default python-pptx name, NOT an opaque id
    buf = io.BytesIO()
    prs.save(buf)
    slide_b64 = base64.b64encode(buf.getvalue()).decode("ascii")

    upd = ScatterChartUpdate(
        chartType="Scatter",
        series=[ScatterSeries(name="p", points=[ScatterPoint(x=1, y=2), ScatterPoint(x=3, y=4)])],
    )
    result = await chart_engine.update_chart_in_slide_base64(slide_b64, legacy_eid, upd)
    new_eid = result["elementId"]
    assert new_eid.startswith("oasp-chart-")  # migrated off the legacy form
    assert new_eid != legacy_eid

    # The recreated chart now carries the opaque id in its OOXML name and is located by it.
    got = await chart_engine.get_chart_from_slide_base64(result["slideBase64"], new_eid)
    assert got["chart"]["chartType"] == "Scatter"
    prs_reloaded = Presentation(io.BytesIO(base64.b64decode(result["slideBase64"])))
    names = [s.name for sl in prs_reloaded.slides for s in sl.shapes if getattr(s, "has_chart", False)]
    assert new_eid in names


# ---------------------------------------------------------------------------
# Path B — in-memory base64 single-slide packages — #13
# ---------------------------------------------------------------------------


def _blank_slide_base64() -> str:
    """A base64-encoded single-slide .pptx (mimics the Add-In's Slide.exportAsBase64)."""
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    buf = io.BytesIO()
    prs.save(buf)
    return base64.b64encode(buf.getvalue()).decode("ascii")


@pytest.mark.asyncio
async def test_generate_standalone_slide_base64_round_trip() -> None:
    """生成独立单页: in-memory deck → base64 → load → parse back to ChartData (no disk)."""
    chart = CategoricalChartData(
        chartType="ColumnClustered",
        categories=["Q1", "Q2", "Q3"],
        series=[CategoricalSeries(name="Rev", values=[10, 20, 30])],
        title="Quarterly",
        showLegend=True,
    )
    result = await chart_engine.generate_chart_slide_base64(chart, ChartInsertOptions(width=400, height=300))
    assert result["elementId"].startswith("oasp-chart-")
    assert result["slideIndex"] == 0
    assert result["chartType"] == "ColumnClustered"
    assert result["seriesCount"] == 1
    b64 = result["slideBase64"]

    # The package is a valid, self-contained 1-slide deck.
    prs = Presentation(io.BytesIO(base64.b64decode(b64)))
    assert len(prs.slides) == 1

    # Read the chart back out of the package by its opaque id.
    got = await chart_engine.get_chart_from_slide_base64(b64, result["elementId"])
    assert got["elementId"] == result["elementId"]
    assert got["chart"]["chartType"] == "ColumnClustered"
    assert got["chart"]["categories"] == ["Q1", "Q2", "Q3"]
    assert got["chart"]["title"] == "Quarterly"
    assert got["chart"]["series"][0]["values"] == [10.0, 20.0, 30.0]


@pytest.mark.asyncio
async def test_generate_standalone_scatter_slide_base64() -> None:
    chart = ScatterChartData(
        chartType="Scatter",
        series=[ScatterSeries(name="ads", points=[ScatterPoint(x=1, y=10), ScatterPoint(x=2, y=20)])],
        title="xy",
    )
    result = await chart_engine.generate_chart_slide_base64(chart)
    got = await chart_engine.get_chart_from_slide_base64(result["slideBase64"], result["elementId"])
    assert got["chart"]["chartType"] == "Scatter"
    assert got["chart"]["series"][0]["points"] == [{"x": 1.0, "y": 10.0}, {"x": 2.0, "y": 20.0}]


@pytest.mark.asyncio
async def test_insert_chart_into_existing_slide_base64() -> None:
    """改单页(插现有页): load exported slide base64 → add chart → back to base64."""
    slide_b64 = _blank_slide_base64()
    chart = CategoricalChartData(
        chartType="Pie",
        categories=["A", "B"],
        series=[CategoricalSeries(name="s", values=[3, 7])],
    )
    result = await chart_engine.insert_chart_into_slide_base64(slide_b64, chart, ChartInsertOptions(left=20, top=20))
    assert result["elementId"].startswith("oasp-chart-")
    assert "slideBase64" in result

    got = await chart_engine.get_chart_from_slide_base64(result["slideBase64"], result["elementId"])
    assert got["chart"]["chartType"] == "Pie"
    assert got["chart"]["series"][0]["values"] == [3.0, 7.0]


@pytest.mark.asyncio
async def test_update_chart_in_slide_base64_round_trip() -> None:
    """改单页(改图): generate → update title in base64 → re-read; opaque id stays stable."""
    chart = CategoricalChartData(
        chartType="BarClustered",
        categories=["A", "B"],
        series=[CategoricalSeries(name="s", values=[1, 2])],
        title="Old",
    )
    gen = await chart_engine.generate_chart_slide_base64(chart)
    eid = gen["elementId"]

    upd = CategoricalChartUpdate(chartType="BarClustered", title="New")
    result = await chart_engine.update_chart_in_slide_base64(gen["slideBase64"], eid, upd)
    assert result["elementId"] == eid
    assert "title" in result["updatedFields"]

    got = await chart_engine.get_chart_from_slide_base64(result["slideBase64"], eid)
    assert got["chart"]["title"] == "New"


@pytest.mark.asyncio
async def test_update_data_in_slide_base64() -> None:
    chart = CategoricalChartData(
        chartType="ColumnClustered",
        categories=["A", "B"],
        series=[CategoricalSeries(name="s", values=[1, 2])],
    )
    gen = await chart_engine.generate_chart_slide_base64(chart)
    upd = CategoricalChartUpdate(
        chartType="ColumnClustered",
        categories=["X", "Y", "Z"],
        series=[CategoricalSeries(name="t", values=[10, 20, 30])],
    )
    result = await chart_engine.update_chart_in_slide_base64(gen["slideBase64"], gen["elementId"], upd)
    got = await chart_engine.get_chart_from_slide_base64(result["slideBase64"], gen["elementId"])
    assert got["chart"]["categories"] == ["X", "Y", "Z"]
    assert got["chart"]["series"][0]["values"] == [10.0, 20.0, 30.0]


@pytest.mark.asyncio
async def test_update_cross_variant_in_slide_base64_round_trip() -> None:
    """改单页(改图,跨 variant): base64 deck → recreate Column→Scatter → re-encode → reload locates by same id."""
    chart = CategoricalChartData(
        chartType="ColumnClustered",
        categories=["A", "B"],
        series=[CategoricalSeries(name="s", values=[1, 2])],
    )
    gen = await chart_engine.generate_chart_slide_base64(chart)
    eid = gen["elementId"]

    upd = ScatterChartUpdate(
        chartType="Scatter",
        series=[ScatterSeries(name="p", points=[ScatterPoint(x=1, y=2), ScatterPoint(x=3, y=4)])],
    )
    result = await chart_engine.update_chart_in_slide_base64(gen["slideBase64"], eid, upd)
    assert result["elementId"] == eid  # opaque id preserved across delete-and-recreate in the package
    assert result["chartType"] == "Scatter"

    # The re-encoded package reloads and the same opaque id still locates the (now Scatter) chart.
    got = await chart_engine.get_chart_from_slide_base64(result["slideBase64"], eid)
    assert got["chart"]["chartType"] == "Scatter"
    assert got["chart"]["series"][0]["points"] == [{"x": 1.0, "y": 2.0}, {"x": 3.0, "y": 4.0}]


@pytest.mark.asyncio
async def test_get_from_base64_unknown_element_raises_3010() -> None:
    slide_b64 = _blank_slide_base64()
    with pytest.raises(ChartEngineError) as exc:
        await chart_engine.get_chart_from_slide_base64(slide_b64, "oasp-chart-does-not-exist")
    assert exc.value.code == ErrorCode.ELEMENT_NOT_FOUND


@pytest.mark.asyncio
async def test_load_malformed_base64_raises_4002() -> None:
    with pytest.raises(ChartEngineError) as exc:
        await chart_engine.get_chart_from_slide_base64("not-valid-base64-@@@", "oasp-chart-x")
    assert exc.value.code == ErrorCode.INVALID_PARAM


@pytest.mark.asyncio
async def test_load_valid_base64_but_not_pptx_raises_4002() -> None:
    junk = base64.b64encode(b"this is not a pptx package").decode("ascii")
    with pytest.raises(ChartEngineError) as exc:
        await chart_engine.get_chart_from_slide_base64(junk, "oasp-chart-x")
    assert exc.value.code == ErrorCode.INVALID_PARAM


@pytest.mark.asyncio
async def test_insert_into_base64_with_no_slides_raises_4002() -> None:
    """A package with zero slides cannot host a chart-into-existing-page insert."""
    prs = Presentation()  # default template has zero slides
    buf = io.BytesIO()
    prs.save(buf)
    empty_b64 = base64.b64encode(buf.getvalue()).decode("ascii")
    chart = CategoricalChartData(
        chartType="Line",
        categories=["A"],
        series=[CategoricalSeries(name="s", values=[1])],
    )
    with pytest.raises(ChartEngineError) as exc:
        await chart_engine.insert_chart_into_slide_base64(empty_b64, chart, None)
    assert exc.value.code == ErrorCode.INVALID_PARAM


# ---------------------------------------------------------------------------
# title=None deletion / partial-update preservation / display options
# (subtle model_fields_set semantics — easy to silently break on a DTO change)
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_update_title_explicit_none_deletes_title() -> None:
    """Explicit ``title=None`` (present in model_fields_set) removes the chart title."""
    chart = CategoricalChartData(
        chartType="ColumnClustered",
        categories=["A", "B"],
        series=[CategoricalSeries(name="s", values=[1, 2])],
        title="Original",
    )
    gen = await chart_engine.generate_chart_slide_base64(chart)
    eid = gen["elementId"]

    # model_validate guarantees "title" lands in model_fields_set even though the value is None.
    upd = CategoricalChartUpdate.model_validate({"chartType": "ColumnClustered", "title": None})
    result = await chart_engine.update_chart_in_slide_base64(gen["slideBase64"], eid, upd)
    assert "title" in result["updatedFields"]

    got = await chart_engine.get_chart_from_slide_base64(result["slideBase64"], eid)
    assert got["chart"]["title"] is None


@pytest.mark.asyncio
async def test_update_data_only_preserves_existing_title() -> None:
    """A series-only update (title absent from model_fields_set) must NOT wipe the title."""
    chart = CategoricalChartData(
        chartType="ColumnClustered",
        categories=["A", "B"],
        series=[CategoricalSeries(name="s", values=[1, 2])],
        title="Keep Me",
    )
    gen = await chart_engine.generate_chart_slide_base64(chart)
    eid = gen["elementId"]

    upd = CategoricalChartUpdate(chartType="ColumnClustered", series=[CategoricalSeries(name="s", values=[9, 8])])
    result = await chart_engine.update_chart_in_slide_base64(gen["slideBase64"], eid, upd)
    assert "title" not in result["updatedFields"]

    got = await chart_engine.get_chart_from_slide_base64(result["slideBase64"], eid)
    assert got["chart"]["title"] == "Keep Me"
    assert got["chart"]["series"][0]["values"] == [9.0, 8.0]


@pytest.mark.asyncio
async def test_display_options_round_trip() -> None:
    """showLegend / showDataLabels written on generate are read back faithfully."""
    chart = CategoricalChartData(
        chartType="ColumnClustered",
        categories=["A", "B"],
        series=[CategoricalSeries(name="s", values=[1, 2])],
        showLegend=True,
        showDataLabels=True,
    )
    gen = await chart_engine.generate_chart_slide_base64(chart)

    got = await chart_engine.get_chart_from_slide_base64(gen["slideBase64"], gen["elementId"])
    assert got["chart"]["showLegend"] is True
    assert got["chart"]["showDataLabels"] is True


@pytest.mark.asyncio
async def test_display_options_legend_off_round_trip() -> None:
    """showLegend=False is honoured (not just the True case)."""
    chart = CategoricalChartData(
        chartType="Pie",
        categories=["A", "B"],
        series=[CategoricalSeries(name="s", values=[1, 2])],
        showLegend=False,
    )
    gen = await chart_engine.generate_chart_slide_base64(chart)

    got = await chart_engine.get_chart_from_slide_base64(gen["slideBase64"], gen["elementId"])
    assert got["chart"]["showLegend"] is False
