"""instantiate_from_template / find_soffice 单测 | Unit tests for template instantiation."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from office4ai.office.authoring.helpers import (
    SofficeNotFoundError,
    TemplateFormatError,
    find_soffice,
    instantiate_from_template,
)
from office4ai.office.authoring.helpers._ooxml import CT_POTX, CT_PPTX, read_package
from office4ai.office.authoring.helpers._soffice import SOFFICE_ENV_VAR


class TestInstantiatePotx:
    def test_swaps_content_type_and_opens(self, potx_template: Path, tmp_path: Path) -> None:
        out = instantiate_from_template(potx_template, tmp_path / "deck.pptx")
        assert out.exists()
        Presentation(out)  # 必须能作为合法演示文稿打开
        ct = read_package(out)[1]["[Content_Types].xml"]
        assert CT_PPTX.encode() in ct and CT_POTX.encode() not in ct

    def test_creates_output_parent_dirs(self, potx_template: Path, tmp_path: Path) -> None:
        out = instantiate_from_template(potx_template, tmp_path / "nested" / "deck.pptx")
        assert out.exists()

    def test_non_potx_disguised_raises(self, tmp_path: Path) -> None:
        # 一份普通 pptx 改名成 .potx —— 缺 template content-type，应报错
        fake = tmp_path / "fake.potx"
        Presentation().save(fake)
        with pytest.raises(TemplateFormatError):
            instantiate_from_template(fake, tmp_path / "o.pptx")


class TestInstantiateUnsupported:
    def test_unsupported_extension_raises(self, tmp_path: Path) -> None:
        doc = tmp_path / "plain.docx"
        Document().save(doc)
        with pytest.raises(TemplateFormatError):
            instantiate_from_template(doc, tmp_path / "o.docx")


class TestInstantiateViaSoffice:
    def test_dotx_to_docx(self, dotx_template: Path, tmp_path: Path, soffice_or_skip: str) -> None:
        out = instantiate_from_template(dotx_template, tmp_path / "doc.docx")
        assert out.exists()
        Document(out)  # 实例化结果须为合法文档

    def test_xltx_to_xlsx(self, xltx_template: Path, tmp_path: Path, soffice_or_skip: str) -> None:
        out = instantiate_from_template(xltx_template, tmp_path / "book.xlsx")
        assert out.exists()
        load_workbook(out)  # 实例化结果须为合法工作簿


class TestFindSoffice:
    def test_env_override_wins(self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
        fake = tmp_path / "soffice"
        fake.write_text("#!/bin/sh\n")
        monkeypatch.setenv(SOFFICE_ENV_VAR, str(fake))
        assert find_soffice() == str(fake)

    def test_explicit_arg_wins(self, tmp_path: Path) -> None:
        fake = tmp_path / "soffice"
        fake.write_text("#!/bin/sh\n")
        assert find_soffice(str(fake)) == str(fake)

    def test_not_found_raises(self, monkeypatch: pytest.MonkeyPatch) -> None:
        monkeypatch.delenv(SOFFICE_ENV_VAR, raising=False)
        monkeypatch.setattr("shutil.which", lambda _name: None)
        monkeypatch.setattr("office4ai.office.authoring.helpers._soffice._COMMON_SOFFICE_PATHS", ())
        with pytest.raises(SofficeNotFoundError):
            find_soffice()
