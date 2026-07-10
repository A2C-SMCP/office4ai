"""fill_sdt_controls / locate_by_style / wrap_in_sdt / replace_runs_with_token 单测。

Unit tests for the Word primitives（含 W3 抽取方向的 wrap_in_sdt / replace_runs_with_token）。
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from office4ai.office.authoring.helpers import (
    AnchorNotFoundError,
    SdtCountMismatchError,
    fill_sdt_controls,
    locate_by_style,
    replace_runs_with_token,
    wrap_in_sdt,
)
from office4ai.office.authoring.helpers._ooxml import parse_xml, qn, read_package


def _document_xml(docx: Path) -> str:
    _, items = read_package(docx)
    return items["word/document.xml"].decode("utf-8")


def _sdt_texts(docx: Path) -> list[str]:
    """按文档顺序取每个 SDT 的可见文本，用于断言填充落点。"""
    _, items = read_package(docx)
    doc = parse_xml(items["word/document.xml"])
    out: list[str] = []
    for sdt in doc.iter(qn("w:sdt")):
        content = sdt.find(qn("w:sdtContent"))
        if content is None:
            continue
        texts = list(content.iter(qn("w:t")))
        if texts:
            out.append("".join(t.text or "" for t in texts))
    return out


class TestFillSdtControls:
    def test_list_mode_fills_in_order(self, sdt_docx: Path, tmp_path: Path) -> None:
        out = tmp_path / "out.docx"
        n = fill_sdt_controls(sdt_docx, ["标题A", "副标题B", "章节C"], output=out)
        assert n == 3
        assert _sdt_texts(out) == ["标题A", "副标题B", "章节C"]

    def test_dict_mode_matches_by_alias(self, sdt_docx: Path, tmp_path: Path) -> None:
        out = tmp_path / "out.docx"
        n = fill_sdt_controls(
            sdt_docx,
            {"文档标题": "2026 报告", "副标题": "副标题内容", "章节一标题": "第一章"},
            output=out,
        )
        assert n == 3
        assert _sdt_texts(out) == ["2026 报告", "副标题内容", "第一章"]

    def test_removes_placeholder_marks(self, sdt_docx: Path, tmp_path: Path) -> None:
        out = tmp_path / "out.docx"
        fill_sdt_controls(sdt_docx, ["a", "b", "c"], output=out)
        xml = _document_xml(out)
        assert "showingPlcHdr" not in xml
        assert 'xml:space="preserve"' in xml  # 首尾空格保护标记已加

    def test_in_place_edit(self, sdt_docx: Path) -> None:
        n = fill_sdt_controls(sdt_docx, ["x", "y", "z"])  # output=None -> 原地
        assert n == 3
        assert _sdt_texts(sdt_docx) == ["x", "y", "z"]

    def test_strict_count_mismatch_raises(self, sdt_docx: Path) -> None:
        with pytest.raises(SdtCountMismatchError):
            fill_sdt_controls(sdt_docx, ["only-one"])

    def test_non_strict_count_mismatch_allowed(self, sdt_docx: Path, tmp_path: Path) -> None:
        out = tmp_path / "out.docx"
        n = fill_sdt_controls(sdt_docx, ["one", "two"], output=out, strict=False)
        assert n == 2  # 只填了前两个

    def test_dict_unknown_alias_strict_raises(self, sdt_docx: Path) -> None:
        with pytest.raises(AnchorNotFoundError):
            fill_sdt_controls(sdt_docx, {"不存在的alias": "v"})

    def test_dict_unknown_alias_non_strict_ok(self, sdt_docx: Path, tmp_path: Path) -> None:
        out = tmp_path / "out.docx"
        n = fill_sdt_controls(sdt_docx, {"文档标题": "T", "缺失": "x"}, output=out, strict=False)
        assert n == 1

    def test_rejects_str_values(self, sdt_docx: Path) -> None:
        with pytest.raises(TypeError):
            fill_sdt_controls(sdt_docx, "整串会被逐字符拆开")  # type: ignore[arg-type]


class TestLocateByStyle:
    def test_finds_heading_paragraphs(self, sdt_docx: Path) -> None:
        matches = locate_by_style(sdt_docx, "Heading1")
        texts = [m.text for m in matches]
        assert "普通章节正文" in texts
        assert all(isinstance(m.index, int) for m in matches)

    def test_unknown_style_returns_empty(self, sdt_docx: Path) -> None:
        assert locate_by_style(sdt_docx, "不存在的样式") == []


# ── W3 抽取方向的夹具与断言工具 | fixtures & asserts for the extract direction ──


def _build_docx(path: Path, paragraphs: list[list[str]], *, styles: dict[int, str] | None = None) -> Path:
    """按「每段一组 run 文本」程序化建 .docx；``styles`` 给指定段落设样式（模拟 run-splitting）。"""
    styles = styles or {}
    doc = Document()
    for i, runs in enumerate(paragraphs):
        para = doc.add_paragraph()
        if i in styles:
            para.style = styles[i]
        for text in runs:
            para.add_run(text)
    doc.save(str(path))
    return path


def _para_t_texts(docx: Path) -> list[list[str]]:
    """取每个正文段落的 w:t 文本列表（保留 run 边界），用于断言 token 是否落在单个 run。"""
    _, items = read_package(docx)
    doc = parse_xml(items["word/document.xml"])
    body = doc.find(qn("w:body"))
    out: list[list[str]] = []
    if body is None:
        return out
    for para in body.iter(qn("w:p")):
        out.append([t.text or "" for t in para.iter(qn("w:t"))])
    return out


class TestWrapInSdt:
    def test_wraps_paragraph_and_fill_roundtrip(self, tmp_path: Path) -> None:
        src = _build_docx(tmp_path / "ref.docx", [["报告标题实值"], ["正文"]])
        out = tmp_path / "tpl.docx"
        wrap_in_sdt(src, 0, "标题", placeholder="【标题】", output=out)

        xml = _document_xml(out)
        assert "w:sdt" in xml and "showingPlcHdr" in xml
        assert _sdt_texts(out) == ["【标题】"]  # 占位文本已归一

        # 抽取(wrap) ↔ 填充(fill) 闭环：产出的模板可被 fill_sdt_controls 按 alias 回填
        filled = tmp_path / "filled.docx"
        n = fill_sdt_controls(out, {"标题": "2026 年度报告"}, output=filled)
        assert n == 1
        assert _sdt_texts(filled) == ["2026 年度报告"]
        assert "showingPlcHdr" not in _document_xml(filled)  # 回填后灰显标记被剥离

    def test_locate_then_wrap_by_index(self, tmp_path: Path) -> None:
        src = _build_docx(
            tmp_path / "ref.docx",
            [["前言"], ["第一章标题"], ["正文"]],
            styles={1: "Heading 1"},
        )
        (match,) = locate_by_style(src, "Heading1")
        out = tmp_path / "tpl.docx"
        wrap_in_sdt(src, match.index, "章标题", output=out)  # 定位坐标直接喂 wrap
        # 被样式命中的那段现已包进 SDT，且 alias 命中可回填
        n = fill_sdt_controls(out, {"章标题": "总体情况"}, output=tmp_path / "f.docx")
        assert n == 1

    def test_no_placeholder_keeps_original_text(self, tmp_path: Path) -> None:
        src = _build_docx(tmp_path / "ref.docx", [["原始标题"]])
        out = tmp_path / "tpl.docx"
        wrap_in_sdt(src, 0, "标题", output=out)
        assert _sdt_texts(out) == ["原始标题"]  # placeholder=None 保留原文本

    def test_tag_defaults_to_alias(self, tmp_path: Path) -> None:
        src = _build_docx(tmp_path / "ref.docx", [["X"]])
        out = tmp_path / "tpl.docx"
        wrap_in_sdt(src, 0, "甲", output=out)
        xml = _document_xml(out)
        assert 'w:val="甲"' in xml  # alias 与 tag 均为「甲」

    def test_sdt_id_passthrough(self, tmp_path: Path) -> None:
        src = _build_docx(tmp_path / "ref.docx", [["X"]])
        out = tmp_path / "tpl.docx"
        wrap_in_sdt(src, 0, "标题", sdt_id=4242, output=out)
        assert 'w:val="4242"' in _document_xml(out)  # 显式 sdt_id 落进 w:id

    def test_empty_paragraph_gets_placeholder_run(self, tmp_path: Path) -> None:
        # 空段落（无 run）+ placeholder：应补一个 run 承载占位文本
        src = _build_docx(tmp_path / "ref.docx", [[]])
        out = tmp_path / "tpl.docx"
        wrap_in_sdt(src, 0, "标题", placeholder="【占位】", output=out)
        assert _sdt_texts(out) == ["【占位】"]

    def test_out_of_range_raises(self, tmp_path: Path) -> None:
        src = _build_docx(tmp_path / "ref.docx", [["only"]])
        with pytest.raises(AnchorNotFoundError):
            wrap_in_sdt(src, 5, "x", output=tmp_path / "o.docx")


class TestReplaceRunsWithToken:
    def test_single_run_replace(self, tmp_path: Path) -> None:
        src = _build_docx(tmp_path / "ref.docx", [["尊敬的张三先生"]])
        out = tmp_path / "tpl.docx"
        n = replace_runs_with_token(src, {"张三": "{{name}}"}, output=out)
        assert n == 1
        assert "".join(_para_t_texts(out)[0]) == "尊敬的{{name}}先生"

    def test_replace_across_split_runs_keeps_token_contiguous(self, tmp_path: Path) -> None:
        # 实值「2026 年度报告」被拆进多个 run —— 模拟 Word 的 run-splitting
        src = _build_docx(tmp_path / "ref.docx", [["标题：", "2026", " 年度", "报告", "（终稿）"]])
        out = tmp_path / "tpl.docx"
        n = replace_runs_with_token(src, {"2026 年度报告": "{{title}}"}, output=out)
        assert n == 1
        runs = _para_t_texts(out)[0]
        assert "".join(runs) == "标题：{{title}}（终稿）"
        # 关键：{{title}} 必须整体落在**单个** run，否则 docxtpl 认不出
        assert any("{{title}}" in r for r in runs), f"token 被拆断: {runs}"

    def test_multiple_occurrences_counted(self, tmp_path: Path) -> None:
        src = _build_docx(tmp_path / "ref.docx", [["甲张三乙", "张三丙"]])
        out = tmp_path / "tpl.docx"
        n = replace_runs_with_token(src, {"张三": "{{n}}"}, output=out)
        assert n == 2
        assert "".join(_para_t_texts(out)[0]) == "甲{{n}}乙{{n}}丙"

    def test_multiple_tokens_in_one_paragraph(self, tmp_path: Path) -> None:
        src = _build_docx(tmp_path / "ref.docx", [["张三于2026年签署"]])
        out = tmp_path / "tpl.docx"
        n = replace_runs_with_token(src, {"张三": "{{name}}", "2026": "{{year}}"}, output=out)
        assert n == 2
        assert "".join(_para_t_texts(out)[0]) == "{{name}}于{{year}}年签署"

    def test_strict_miss_raises(self, tmp_path: Path) -> None:
        src = _build_docx(tmp_path / "ref.docx", [["无关内容"]])
        with pytest.raises(AnchorNotFoundError):
            replace_runs_with_token(src, {"缺失实值": "{{x}}"}, output=tmp_path / "o.docx")

    def test_non_strict_miss_allowed(self, tmp_path: Path) -> None:
        src = _build_docx(tmp_path / "ref.docx", [["有张三"]])
        out = tmp_path / "tpl.docx"
        n = replace_runs_with_token(src, {"张三": "{{a}}", "缺失": "{{b}}"}, output=out, strict=False)
        assert n == 1
