"""create_custom_layout / locate_by_master / promote_slide_to_layout 单测。

Unit tests for the PPT primitives（含 W3 抽取方向的 promote_slide_to_layout）。
"""

from __future__ import annotations

from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Emu

from office4ai.office.authoring.helpers import (
    AnchorNotFoundError,
    ExternalMediaError,
    create_custom_layout,
    locate_by_master,
    promote_slide_to_layout,
)

LAYOUT_NAME = "霓虹内容页"


def _decorate(shapes) -> None:
    """在临时页画一个带标记文字的文本框，便于断言装饰进入了版式。"""
    box = shapes.add_textbox(Emu(914400), Emu(914400), Emu(3000000), Emu(500000))
    box.text_frame.paragraphs[0].add_run().text = "BRAND_MARK"


def _add_marked_slide(prs: Presentation, mark: str = "SAMPLE_MARK") -> None:
    """往 deck 加一张带标记文本框的样板页（供 promote_slide_to_layout 抽取）。"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # 空白版式
    box = slide.shapes.add_textbox(Emu(914400), Emu(914400), Emu(3000000), Emu(500000))
    box.text_frame.paragraphs[0].add_run().text = mark


def _layout_text(layout) -> str:
    return "".join(
        run.text
        for shape in layout.shapes
        if shape.has_text_frame
        for para in shape.text_frame.paragraphs
        for run in para.runs
    )


class TestCreateCustomLayout:
    def test_registers_layout(self, blank_prs: Presentation) -> None:
        before = len(blank_prs.slide_masters[0].slide_layouts)
        layout = create_custom_layout(blank_prs, LAYOUT_NAME, _decorate)
        assert layout.name == LAYOUT_NAME
        assert len(blank_prs.slide_masters[0].slide_layouts) == before + 1

    def test_removes_temp_slide(self, blank_prs: Presentation) -> None:
        create_custom_layout(blank_prs, LAYOUT_NAME, _decorate)
        assert len(blank_prs.slides) == 0, "临时页应被删除，成品不应多一页"

    def test_decoration_lands_in_layout(self, blank_prs: Presentation, tmp_path: Path) -> None:
        create_custom_layout(blank_prs, LAYOUT_NAME, _decorate)
        out = tmp_path / "deck.pptx"
        blank_prs.save(out)
        layout = locate_by_master(Presentation(out), LAYOUT_NAME)
        text = "".join(
            run.text
            for shape in layout.shapes
            if shape.has_text_frame
            for para in shape.text_frame.paragraphs
            for run in para.runs
        )
        assert "BRAND_MARK" in text

    def test_survives_roundtrip_and_is_usable(self, blank_prs: Presentation, tmp_path: Path) -> None:
        create_custom_layout(blank_prs, LAYOUT_NAME, _decorate)
        out = tmp_path / "deck.pptx"
        blank_prs.save(out)
        reloaded = Presentation(out)
        layout = locate_by_master(reloaded, LAYOUT_NAME)
        reloaded.slides.add_slide(layout)  # 复用版式加一页，不应抛错
        used = tmp_path / "used.pptx"
        reloaded.save(used)
        assert len(Presentation(used).slides) == 1

    def test_bad_base_index_raises(self, blank_prs: Presentation) -> None:
        with pytest.raises(ValueError, match="out of range|越界"):
            create_custom_layout(blank_prs, LAYOUT_NAME, _decorate, base_layout_index=999)

    def test_media_decoration_fails_fast(self, blank_prs: Presentation, tmp_path: Path) -> None:
        # 装饰若加了图片，其 r:embed 不随形状搬运 → 同样应快速失败而非静默产出损坏版式
        img = tmp_path / "logo.png"
        Image.new("RGB", (2, 2), "blue").save(img)

        def _decorate_with_image(shapes) -> None:
            shapes.add_picture(str(img), Emu(0), Emu(0))

        with pytest.raises(ExternalMediaError):
            create_custom_layout(blank_prs, LAYOUT_NAME, _decorate_with_image)


class TestLocateByMaster:
    def test_found_after_create(self, blank_prs: Presentation) -> None:
        create_custom_layout(blank_prs, LAYOUT_NAME, _decorate)
        assert locate_by_master(blank_prs, LAYOUT_NAME).name == LAYOUT_NAME

    def test_not_found_raises(self, blank_prs: Presentation) -> None:
        with pytest.raises(AnchorNotFoundError):
            locate_by_master(blank_prs, "不存在的版式")


PROMOTED_NAME = "样板版式"


class TestPromoteSlideToLayout:
    def test_registers_layout_from_existing_slide(self, blank_prs: Presentation) -> None:
        _add_marked_slide(blank_prs)
        before = len(blank_prs.slide_masters[0].slide_layouts)
        layout = promote_slide_to_layout(blank_prs, 0, PROMOTED_NAME)
        assert layout.name == PROMOTED_NAME
        assert len(blank_prs.slide_masters[0].slide_layouts) == before + 1

    def test_source_slide_retained_by_default(self, blank_prs: Presentation) -> None:
        _add_marked_slide(blank_prs)
        promote_slide_to_layout(blank_prs, 0, PROMOTED_NAME)
        assert len(blank_prs.slides) == 1  # 默认保留源页

    def test_remove_source_drops_the_slide(self, blank_prs: Presentation) -> None:
        _add_marked_slide(blank_prs)
        promote_slide_to_layout(blank_prs, 0, PROMOTED_NAME, remove_source=True)
        assert len(blank_prs.slides) == 0  # 产出干净模板

    def test_shapes_land_in_layout_after_roundtrip(self, blank_prs: Presentation, tmp_path: Path) -> None:
        _add_marked_slide(blank_prs)
        promote_slide_to_layout(blank_prs, 0, PROMOTED_NAME, remove_source=True)
        out = tmp_path / "tpl.pptx"
        blank_prs.save(out)
        layout = locate_by_master(Presentation(out), PROMOTED_NAME)
        assert "SAMPLE_MARK" in _layout_text(layout)

    def test_promoted_layout_is_reusable(self, blank_prs: Presentation, tmp_path: Path) -> None:
        _add_marked_slide(blank_prs)
        promote_slide_to_layout(blank_prs, 0, PROMOTED_NAME, remove_source=True)
        out = tmp_path / "tpl.pptx"
        blank_prs.save(out)
        reloaded = Presentation(out)
        layout = locate_by_master(reloaded, PROMOTED_NAME)
        reloaded.slides.add_slide(layout)  # 用抽出的版式加页，不应抛错
        used = tmp_path / "used.pptx"
        reloaded.save(used)
        assert len(Presentation(used).slides) == 1

    def test_bad_slide_index_raises(self, blank_prs: Presentation) -> None:
        _add_marked_slide(blank_prs)
        with pytest.raises(ValueError, match="out of range|越界"):
            promote_slide_to_layout(blank_prs, 9, PROMOTED_NAME)

    def test_bad_base_index_raises(self, blank_prs: Presentation) -> None:
        _add_marked_slide(blank_prs)
        with pytest.raises(ValueError, match="out of range|越界"):
            promote_slide_to_layout(blank_prs, 0, PROMOTED_NAME, base_layout_index=999)

    def test_media_sample_slide_fails_fast(self, blank_prs: Presentation, tmp_path: Path) -> None:
        # 含图片的样板页：其 r:embed 引用不随形状搬运，抽版式会静默产出损坏文件 → 应快速失败
        img = tmp_path / "logo.png"
        Image.new("RGB", (2, 2), "red").save(img)
        slide = blank_prs.slides.add_slide(blank_prs.slide_layouts[6])
        slide.shapes.add_picture(str(img), Emu(0), Emu(0))
        with pytest.raises(ExternalMediaError):
            promote_slide_to_layout(blank_prs, 0, PROMOTED_NAME)
