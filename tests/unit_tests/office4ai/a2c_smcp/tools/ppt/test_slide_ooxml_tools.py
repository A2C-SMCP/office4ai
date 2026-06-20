"""Whole-slide OOXML 工具单元测试 (office4ai #40).

测试策略 (对齐 test_ppt_tools.py):
- Mock OfficeWorkspace.execute() 的返回值
- 验证 source_path → base64 翻译 (整页 base64 进 wire，不进入参)
- 验证 OfficeAction 构建 (category / event_name / camelCase params)
- 验证导出 base64 落盘到 dest_path、返回句柄不含 base64
- 验证错误路径 (文件缺失/空文件/业务失败)
"""

from __future__ import annotations

import base64
from unittest.mock import AsyncMock, MagicMock

import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

from office4ai.a2c_smcp.tools.ppt import PptGetSlideOoxmlTool, PptInsertSlidesOoxmlTool
from office4ai.environment.workspace.base import OfficeObs


@pytest.fixture
def mock_workspace():
    workspace = MagicMock()
    workspace.execute = AsyncMock()
    return workspace


def _make_pptx(path) -> bytes:
    """用 python-pptx 造一个最小但真实的单页 .pptx，返回其字节。"""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text_frame.text = "High fidelity via OOXML"
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    prs.save(str(path))
    return path.read_bytes()


# ============================================================================
# Metadata
# ============================================================================


def test_metadata(mock_workspace):
    insert = PptInsertSlidesOoxmlTool(mock_workspace)
    assert insert.name == "ppt_insert_slides_ooxml"
    assert insert.category == "ppt"
    assert insert.event_name == "insert:slidesOoxml"

    get = PptGetSlideOoxmlTool(mock_workspace)
    assert get.name == "ppt_get_slide_ooxml"
    assert get.category == "ppt"
    assert get.event_name == "get:slideOoxml"


# ============================================================================
# Insert
# ============================================================================


@pytest.mark.asyncio
async def test_insert_reads_file_and_sends_base64(mock_workspace, tmp_path):
    """source_path 被读出、base64 编码后作为 wire base64 发送 (而非 source_path)。"""
    pptx_path = tmp_path / "slide.pptx"
    raw = _make_pptx(pptx_path)

    mock_workspace.execute.return_value = OfficeObs(
        success=True,
        data={"insertedSlideIndices": [1], "insertedSlideIds": ["sid-1"], "elements": []},
    )

    tool = PptInsertSlidesOoxmlTool(mock_workspace)
    result = await tool.execute(
        {
            "document_uri": "file:///tmp/deck.pptx",
            "source_path": str(pptx_path),
            "targetSlideIndex": 0,
            "formatting": "keepSourceFormatting",
        }
    )

    assert result["success"] is True
    assert result["data"]["insertedSlideIds"] == ["sid-1"]

    # 校验发送的 OfficeAction
    action = mock_workspace.execute.call_args.args[0]
    assert action.category == "ppt"
    assert action.action_name == "insert:slidesOoxml"
    params = action.params
    assert params["document_uri"] == "file:///tmp/deck.pptx"
    assert "source_path" not in params  # 路径不上 wire
    assert params["targetSlideIndex"] == 0
    assert params["formatting"] == "keepSourceFormatting"
    # base64 正确且可解码回原字节
    assert base64.b64decode(params["base64"]) == raw


@pytest.mark.asyncio
async def test_insert_missing_file(mock_workspace):
    tool = PptInsertSlidesOoxmlTool(mock_workspace)
    result = await tool.execute({"document_uri": "file:///tmp/deck.pptx", "source_path": "/no/such/file.pptx"})
    assert result["success"] is False
    assert "source_path" in result["error"]
    mock_workspace.execute.assert_not_called()


@pytest.mark.asyncio
async def test_insert_empty_file(mock_workspace, tmp_path):
    empty = tmp_path / "empty.pptx"
    empty.write_bytes(b"")
    tool = PptInsertSlidesOoxmlTool(mock_workspace)
    result = await tool.execute({"document_uri": "file:///tmp/deck.pptx", "source_path": str(empty)})
    assert result["success"] is False
    assert "empty" in result["error"]
    mock_workspace.execute.assert_not_called()


@pytest.mark.asyncio
async def test_insert_business_failure_propagates(mock_workspace, tmp_path):
    pptx_path = tmp_path / "slide.pptx"
    _make_pptx(pptx_path)
    mock_workspace.execute.return_value = OfficeObs(success=False, data={}, error="3003: Document not connected")
    tool = PptInsertSlidesOoxmlTool(mock_workspace)
    result = await tool.execute({"document_uri": "file:///tmp/deck.pptx", "source_path": str(pptx_path)})
    assert result["success"] is False
    assert "3003" in result["error"]


@pytest.mark.asyncio
async def test_insert_optional_fields_excluded_when_absent(mock_workspace, tmp_path):
    """未给的可选字段不应出现在 wire params 里。"""
    pptx_path = tmp_path / "slide.pptx"
    _make_pptx(pptx_path)
    mock_workspace.execute.return_value = OfficeObs(success=True, data={"insertedSlideIds": ["x"]})
    tool = PptInsertSlidesOoxmlTool(mock_workspace)
    await tool.execute({"document_uri": "file:///tmp/deck.pptx", "source_path": str(pptx_path)})
    params = mock_workspace.execute.call_args.args[0].params
    assert "targetSlideIndex" not in params
    assert "replaceSlideId" not in params
    assert "finalSlideIndex" not in params
    assert "base64" in params


# ============================================================================
# Get
# ============================================================================


@pytest.mark.asyncio
async def test_get_writes_base64_to_disk_and_returns_handle(mock_workspace, tmp_path):
    """导出的 base64 落盘到 dest_path，返回 {slideId, slideIndex, filePath} 不含 base64。"""
    # 模拟 Add-In 返回的整页 base64 (用真实 pptx 字节)
    sample = tmp_path / "sample.pptx"
    raw = _make_pptx(sample)
    wire_b64 = base64.b64encode(raw).decode("ascii")

    mock_workspace.execute.return_value = OfficeObs(
        success=True, data={"slideIndex": 2, "slideId": "sid-xyz", "base64": wire_b64}
    )

    dest = tmp_path / "exported" / "out.pptx"
    tool = PptGetSlideOoxmlTool(mock_workspace)
    result = await tool.execute({"document_uri": "file:///tmp/deck.pptx", "slideIndex": 2, "dest_path": str(dest)})

    assert result["success"] is True
    assert result["data"]["slideId"] == "sid-xyz"
    assert result["data"]["slideIndex"] == 2
    assert result["data"]["filePath"] == str(dest)
    assert "base64" not in result["data"]  # base64 不回传
    assert dest.read_bytes() == raw  # 已落盘且可被 python-pptx 再读

    # 校验发送事件
    action = mock_workspace.execute.call_args.args[0]
    assert action.action_name == "get:slideOoxml"
    assert action.params["slideIndex"] == 2


@pytest.mark.asyncio
async def test_get_business_failure_propagates(mock_workspace, tmp_path):
    mock_workspace.execute.return_value = OfficeObs(success=False, data={}, error="3003: Document not connected")
    tool = PptGetSlideOoxmlTool(mock_workspace)
    result = await tool.execute(
        {"document_uri": "file:///tmp/deck.pptx", "slideIndex": 0, "dest_path": str(tmp_path / "x.pptx")}
    )
    assert result["success"] is False
    assert "3003" in result["error"]
